import streamlit as st
import pandas as pd
import graphviz
import sqlite3
import glob
import os
import re
import math
import json
import openpyxl
from data_pipeline import process_data, load_saved_mappings, save_new_mapping
from db_manager import init_db, insert_data, execute_query, execute_update
from ai_agent import process_ai_chat  # Tích hợp AI từ file ngoài

# ── YARN PARSER ─────────────────────────────────────────────────
_yarn_import_error = None
try:
    from yarn_parser import init_yarn_table, scan_folder, parse_excel_file, upsert_records
    init_yarn_table()
except Exception as _ye:
    _yarn_import_error = str(_ye)
    def scan_folder(*a, **k): return 0, 0, [f"Lỗi: {_ye}"]
    def parse_excel_file(*a, **k): return []
    def upsert_records(*a, **k): return 0
# ─────────────────────────────────────────────────────────────────

# ── SIZING PIPELINE ───────────────────────────────────────────────
_sizing_import_error = None
try:
    from sizing_pipeline import (
        init_sizing_table, import_sizing_file,
        scan_sizing_folder, scan_sizing_folder_v2,
        detect_and_import, list_files_in_folder,
    )
    init_sizing_table()
except Exception as _se:
    _sizing_import_error = str(_se)
    def import_sizing_file(*a, **k): return {"total": 0, "errors": [str(_se)]}
    def scan_sizing_folder(*a, **k): return {"total_rows": 0, "total_files": 0, "errors": [str(_se)]}
    SHEET_COLS = {}
# ──────────────────────────────────────────────────────────────────

# ── SIZING PIPELINE ───────────────────────────────────────────────
_sizing_import_error = None
try:
    from sizing_pipeline import init_sizing_table, import_sizing_file, scan_sizing_folder
    init_sizing_table()
except Exception as _se:
    _sizing_import_error = str(_se)
    def import_sizing_file(*a, **k): return {"total":0,"errors":[str(_se)]}
    def scan_sizing_folder(*a, **k): return {"total_rows":0,"total_files":0,"errors":[str(_se)]}
# ──────────────────────────────────────────────────────────────────

# ── BEAM INFO ─────────────────────────────────────────────────────
_beam_import_error = None
try:
    from beam_info import init_beam_tables, import_beam_file, get_beam_on_machine
    init_beam_tables()
except Exception as _be:
    _beam_import_error = str(_be)
    def import_beam_file(*a, **k): return {"xuatkho":0,"yccb":0,"total":0,"errors":[str(_be)]}
    def get_beam_on_machine(*a, **k): return {}
# ─────────────────────────────────────────────────────────────────

# --- HÀM CAN THIỆP LÕI DATABASE CHỐNG LỖI ---
def run_db_command(query, params=()):
    db_files = glob.glob("*.db")
    if not db_files:
        for root, dirs, files in os.walk("."):
            for file in files:
                if file.endswith(".db"):
                    db_files.append(os.path.join(root, file))
    
    success = False
    for db in db_files:
        try:
            with sqlite3.connect(db) as conn:
                conn.execute(query, params)
                conn.commit()
                success = True
        except:
            pass
    return success

# Khởi tạo Không gian ảo và Bảng Cấu hình
run_db_command("CREATE TABLE IF NOT EXISTS Yarn_Dictionary (yarn_type TEXT PRIMARY KEY, coefficient REAL)")
run_db_command("CREATE TABLE IF NOT EXISTS AI_Rules (id INTEGER PRIMARY KEY AUTOINCREMENT, rule_text TEXT)")
run_db_command("CREATE TABLE IF NOT EXISTS AutoSync_Config (id INTEGER PRIMARY KEY, folder_path TEXT, cluster_name TEXT, skip_rows INTEGER, keyword_filter TEXT, mapping_json TEXT, template_name TEXT, sheet_name TEXT)")

try: run_db_command("ALTER TABLE AutoSync_Config ADD COLUMN keyword_filter TEXT")
except: pass
try: run_db_command("ALTER TABLE AutoSync_Config ADD COLUMN mapping_json TEXT")
except: pass
try: run_db_command("ALTER TABLE AutoSync_Config ADD COLUMN template_name TEXT")
except: pass
try: run_db_command("ALTER TABLE AutoSync_Config ADD COLUMN sheet_name TEXT")
except: pass

# --- TÌM KIẾM XUYÊN THẤU + NHẬN DIỆN TỪ KHÓA ---
def get_all_excel_csv_files(base_folder, keyword_filter=""):
    all_files = []
    keywords = [k.strip().lower() for k in keyword_filter.split(",")] if keyword_filter else []
    
    if os.path.exists(base_folder):
        for root, dirs, files in os.walk(base_folder):
            for file in files:
                if file.endswith(('.csv', '.xls', '.xlsx')) and not file.startswith('~$'):
                    full_path = os.path.join(root, file)
                    if keywords:
                        if any(k in full_path.lower() for k in keywords):
                            all_files.append(full_path)
                    else:
                        all_files.append(full_path)
    return all_files

# --- NẠP DỮ LIỆU + ĐỌC MULTI-SHEET + TỰ MÒ NĂM + TRÁM NGÀY ---
def run_multi_file_pipeline(u_files, mapping, cluster, s_name, s_rows, auto_sync, is_local):
    try:
        total_rows = 0
        sheet_zone_map = mapping.get('__SHEET_ZONE_MAP__', {})

        for f in u_files:
            fname = os.path.basename(f) if is_local else f.name
            f_path = f if is_local else ""
            
            if not is_local: f.seek(0)
            
            extracted_year = str(pd.Timestamp.now().year) 
            if is_local:
                year_match = re.search(r'\\(20\d{2})\\', f_path)
                if not year_match: year_match = re.search(r'/(20\d{2})/', f_path)
                if year_match: extracted_year = year_match.group(1)

            dfs_to_process = {} 
            try:
                if fname.endswith('.csv'): 
                    dfs_to_process["CSV"] = pd.read_csv(f, skiprows=s_rows)
                else: 
                    xls = pd.ExcelFile(f)
                    if s_name == "ALL_SHEETS":
                        for sheet in xls.sheet_names:
                            dfs_to_process[sheet] = pd.read_excel(xls, sheet_name=sheet, skiprows=s_rows)
                    else:
                        dfs_to_process[s_name] = pd.read_excel(xls, sheet_name=s_name, skiprows=s_rows)
            except Exception as e: 
                continue 
            
            for current_sheet_name, df in dfs_to_process.items():
                if df.empty: continue

                df = df.loc[:, ~df.columns.duplicated()].copy()

                mapped_zone = sheet_zone_map.get(current_sheet_name, current_sheet_name)
                if mapped_zone == "-- Bỏ qua --":
                    continue

                cols_to_keep = []
                rename_dict = {}

                for src, tgt in mapping.items():
                    if src.startswith('__') or src == "AUTO_SHEET_NAME": continue
                    if src in df.columns:
                        if tgt != "-- Bỏ qua --":
                            cols_to_keep.append(src)
                            if tgt != "Giữ nguyên tên cột gốc (Copy từ trái)":
                                rename_dict[src] = tgt

                if not cols_to_keep:
                    continue 

                df_mapped = df[cols_to_keep].copy()
                df_mapped.rename(columns=rename_dict, inplace=True)
                
                df_mapped = df_mapped.loc[:, ~df_mapped.columns.duplicated()]
                
                df_mapped['sub_location'] = mapped_zone
                
                check_cols = [c for c in df_mapped.columns if c != 'sub_location']
                if not check_cols: continue
                df_mapped = df_mapped.dropna(subset=check_cols, how='all')
                if df_mapped.empty: continue
                
                if 'date_month' in df_mapped.columns and 'date_day' in df_mapped.columns:
                    s_year = df_mapped.get('date_year', pd.Series(extracted_year, index=df_mapped.index)).astype(str).replace(r'(?i)none|nan|^$', extracted_year, regex=True)
                    s_month = df_mapped['date_month'].astype(str).replace(r'(?i)none|nan|^$', pd.NA, regex=True).ffill().bfill()
                    s_day = df_mapped['date_day'].astype(str).replace(r'(?i)none|nan|^$', pd.NA, regex=True).ffill().bfill()

                    s_month = s_month.str.extract(r'(\d+)')[0].str.zfill(2)
                    s_day = s_day.str.extract(r'(\d+)')[0].str.zfill(2)

                    valid_mask = s_month.notna() & s_day.notna()
                    df_mapped['date'] = pd.NaT
                    df_mapped.loc[valid_mask, 'date'] = pd.to_datetime(s_year[valid_mask] + '-' + s_month[valid_mask] + '-' + s_day[valid_mask], errors='coerce')
                    df_mapped['date'] = df_mapped['date'].dt.strftime('%Y-%m-%d')
                    df_mapped.drop(columns=['date_month', 'date_day', 'date_year'], errors='ignore', inplace=True)
                elif 'date' in df_mapped.columns:
                    df_mapped['date'] = df_mapped['date'].astype(str).replace(r'(?i)none|nan|^$', pd.NA, regex=True).ffill().bfill()
                    parsed = pd.to_datetime(df_mapped['date'], errors='coerce', dayfirst=True)
                    if parsed.isna().all():
                        parsed = pd.to_datetime(df_mapped['date'], errors='coerce', dayfirst=False)
                    df_mapped['date'] = parsed.dt.strftime('%Y-%m-%d')

                if 'date' not in df_mapped.columns:
                    df_mapped['date'] = pd.Timestamp.now().strftime('%Y-%m-%d')
                else:
                    df_mapped['date'] = df_mapped['date'].astype(str).replace(r'(?i)none|nan|nat|^$', pd.NA, regex=True)
                    if df_mapped['date'].isna().all():
                        df_mapped['date'] = pd.Timestamp.now().strftime('%Y-%m-%d')
                    else:
                        df_mapped['date'] = df_mapped['date'].fillna(pd.Timestamp.now().strftime('%Y-%m-%d'))
                    
                df_mapped['cluster_name'] = cluster
                has_inbound = 'inbound_date' in df_mapped.columns
                has_outbound = 'outbound_date' in df_mapped.columns
                dfs_to_insert = []
                
                if has_inbound:
                    df_in = df_mapped.dropna(subset=['inbound_date']).copy()
                    if not df_in.empty:
                        df_in['date'] = pd.to_datetime(df_in['inbound_date'], errors='coerce').ffill().bfill().dt.strftime('%Y-%m-%d')
                        df_in['type'] = 'NHAP'
                        dfs_to_insert.append(df_in)
                if has_outbound:
                    df_out = df_mapped.dropna(subset=['outbound_date']).copy()
                    if not df_out.empty:
                        df_out['date'] = pd.to_datetime(df_out['outbound_date'], errors='coerce').ffill().bfill().dt.strftime('%Y-%m-%d')
                        df_out['type'] = 'XUAT'
                        dfs_to_insert.append(df_out)
                if not has_inbound and not has_outbound:
                    if 'date' not in df_mapped.columns or df_mapped['date'].isna().all(): 
                        df_mapped['date'] = pd.Timestamp.now().strftime('%Y-%m-%d')
                    if cluster == "Xưởng Dệt": df_mapped['type'] = 'SAN_XUAT'
                    else: df_mapped['type'] = 'NHAP'
                    dfs_to_insert.append(df_mapped)
                    
                for final_df in dfs_to_insert:
                    clean_df = final_df.drop(columns=['inbound_date', 'outbound_date'], errors='ignore')
                    
                    clean_df = clean_df.loc[:, ~clean_df.columns.duplicated()]
                    
                    for col in clean_df.columns:
                        clean_df[col] = clean_df[col].apply(lambda x: "" if pd.isna(x) else str(x))
                    
                    try:
                        db_files = glob.glob("*.db")
                        if db_files:
                            with sqlite3.connect(db_files[0]) as conn:
                                cursor = conn.execute("PRAGMA table_info(Inventory_Log)")
                                existing_cols = [row[1] for row in cursor.fetchall()]
                                for c in clean_df.columns:
                                    if c not in existing_cols:
                                        conn.execute(f"ALTER TABLE Inventory_Log ADD COLUMN '{c}' TEXT")
                                conn.commit()
                    except: pass
                            
                    insert_data(clean_df)
                    total_rows += len(clean_df)

        st.cache_data.clear() 
        return total_rows, len(u_files)
    except Exception as e: 
        return 0, str(e)

# ==========================================
# HÀM LỌC BẢNG DỮ LIỆU THÔNG MINH
# ==========================================
def get_clean_cluster_data(df, cluster_name):
    if df.empty: return df
    df_clean = df.loc[:, ~df.columns.duplicated()].copy()
    
    try:
        valid_cols = ['id', 'date', 'cluster_name', 'sub_location', 'type']
        has_mapping = False
        
        db_files = glob.glob("*.db")
        if db_files:
            with sqlite3.connect(db_files[0]) as conn:
                df_cfg = pd.read_sql_query(f"SELECT mapping_json FROM AutoSync_Config WHERE cluster_name='{cluster_name}' LIMIT 1", conn)
                if not df_cfg.empty and df_cfg.iloc[0]['mapping_json']:
                    mapping_str = str(df_cfg.iloc[0]['mapping_json'])
                    if mapping_str.strip() not in ["", "{}"]:
                        mapping_data = json.loads(mapping_str)
                        has_mapping = True
                        for k, v in mapping_data.items():
                            if k.startswith('__') or k == "AUTO_SHEET_NAME": continue
                            if isinstance(v, str):
                                if v == "Giữ nguyên tên cột gốc (Copy từ trái)":
                                    valid_cols.append(k)
                                elif v != "-- Bỏ qua --":
                                    valid_cols.append(v)
        
        cols_to_keep = []
        for c in df_clean.columns:
            if has_mapping:
                if c in valid_cols:
                    cols_to_keep.append(c)
            else:
                if c in valid_cols:
                    cols_to_keep.append(c)
                else:
                    col_str = df_clean[c].astype(str).str.strip().str.lower()
                    is_empty = col_str.isin(['', 'nan', 'none', '<na>', 'null'])
                    if not is_empty.all():
                        cols_to_keep.append(c)
                        
        safe_cols = []
        seen = set()
        for c in cols_to_keep:
            c_str = str(c)
            if c_str not in seen:
                seen.add(c_str)
                safe_cols.append(c_str)
        cols_to_keep = safe_cols
        
        cols_to_keep = [c for c in cols_to_keep if c in df_clean.columns]
        
        df_clean = df_clean[cols_to_keep].fillna("")
        
        subset_cols = [c for c in df_clean.columns if c != 'id']
        if subset_cols:
            df_clean = df_clean.drop_duplicates(subset=subset_cols)
            
        return df_clean
    except Exception as e:
        return df_clean.fillna("")

# HÀM RADAR TÌM CỘT TỰ ĐỘNG CHO DASHBOARD
def find_best_col(df, keywords):
    for k in keywords:
        for col in df.columns:
            if k in str(col).lower():
                return col
    return None

# 1. THIẾT LẬP TRANG CHÍNH
st.set_page_config(page_title="TexFlow Systems", page_icon="🌊", layout="wide", initial_sidebar_state="expanded")

def _load_css():
    import os
    css_file = os.path.join(os.path.dirname(os.path.abspath(__file__)), "style.css")
    if os.path.exists(css_file):
        with open(css_file, "r", encoding="utf-8") as _f:
            st.markdown(f"<style>{_f.read()}</style>", unsafe_allow_html=True)
    else:
        st.warning("style.css not found")

_load_css()

init_db()

# ── PERSISTENT MEMORY ──────────────────────────────────────────────
_MEMORY_FILE = "chat_memory.json"
_MAX_MEMORY_MSGS = 30

def _load_memory():
    import os
    if os.path.exists(_MEMORY_FILE):
        try:
            with open(_MEMORY_FILE, "r", encoding="utf-8") as _f:
                return json.load(_f).get("messages", [])
        except:
            pass
    return []

def _save_memory(msgs):
    try:
        with open(_MEMORY_FILE, "w", encoding="utf-8") as _f:
            json.dump({"messages": msgs[-_MAX_MEMORY_MSGS:]}, _f, ensure_ascii=False, indent=2)
    except:
        pass
# ───────────────────────────────────────────────────────────────────

# ==========================================
# CƠ CHẾ AUTO-RUN KHI KHỞI ĐỘNG HỆ THỐNG
# ==========================================
# Không tự động sync khi khởi động — người dùng bấm nút thủ công
# ==========================================

# ==========================================
# 2. THANH SIDEBAR
# ==========================================
with st.sidebar:
    st.title("🌊 TexFlow")
    st.caption("Quản lý Chuỗi Cung Ứng Dệt May")
    st.markdown("---")
    menu_selection = st.radio("", ["📊 Dashboard", "📥 Quản lý Dữ liệu"], label_visibility="collapsed")
    st.markdown("---")
    st.subheader("⚙️ Cài đặt Hệ thống AI")
    # Load model_name từ file config (persist qua restart)
    if "model_name" not in st.session_state:
        _saved_model = "qwen2.5:3b"
        for _cfg_path in ["mapping_config.json", "saved_mappings.json"]:
            try:
                if os.path.exists(_cfg_path):
                    with open(_cfg_path, "r", encoding="utf-8") as _fcfg:
                        _cfg_data = json.load(_fcfg)
                        if _cfg_data.get("model_name"):
                            _saved_model = _cfg_data["model_name"]
                            break
            except: pass
        st.session_state.model_name = _saved_model
    # Sidebar: hiện rõ model đang dùng
    st.caption("🤖 Model AI (Ollama):")
    model_name = st.text_input(
        "Tên Model đang chạy trên Ollama:",
        value=st.session_state.model_name,
        key="model_name_input",
    )
    if model_name != st.session_state.model_name:
        st.session_state.model_name = model_name
        # Persist model name to config
        try:
            _cfg_m = {}
            if os.path.exists("mapping_config.json"):
                with open("mapping_config.json","r",encoding="utf-8") as _fm: _cfg_m = json.load(_fm)
            _cfg_m["model_name"] = model_name
            with open("mapping_config.json","w",encoding="utf-8") as _fm2: json.dump(_cfg_m, _fm2)
        except: pass
        # Persist to config file so it survives restarts
        try:
            import json as _jm
            _cfg = {}
            if os.path.exists("mapping_config.json"):
                with open("mapping_config.json","r") as _f: _cfg = json.load(_f)
            _cfg["model_name"] = model_name
            with open("mapping_config.json","w") as _f: json.dump(_cfg, _f)
        except: pass

# ==========================================
# 3. TRANG DASHBOARD & BÁO CÁO TƯƠNG TÁC
# ==========================================
if menu_selection == "📊 Dashboard":
    st.title("TỔNG QUAN QUY TRÌNH DỆT MAY")
    
    graph = graphviz.Digraph(engine='dot')
    graph.attr(rankdir='TB', size='10,8', bgcolor='#1E1E1E') 
    
    node_attr = {
        'shape': 'box', 'style': 'filled,rounded', 'fontname': 'Helvetica', 
        'fontsize': '12', 'fontcolor': '#FFFFFF', 'color': '#4A90E2', 
        'penwidth': '1.5', 'height': '0.9', 'width': '2.2'
    }
    
    d_nodes = {
        "Kho Sợi Tổng": {"color": "#2C3E50", "label": "📦 [KHO SỢI TỔNG]\nNguồn nguyên liệu"},
        "Xưởng Nhuộm": {"color": "#2C3E50", "label": "🎨 [XƯỞNG NHUỘM]\nXử lý màu"},
        "Kho Sợi Sizing": {"color": "#2C3E50", "label": "🧵 [KHO SỢI SIZING]\nBáo cáo số lượng sợi"},
        "Máy Direct": {"color": "#34495E", "label": "⚙️ [MÁY DIRECT]\nSản xuất beam"},
        "Máy Hồ":       {"color": "#34495E", "label": "🧪 [MÁY HỒ]\nHồ sợi"},
        "Máy Winder":   {"color": "#34495E", "label": "🎡 [WINDER]\nQuấn sợi"},
        "Máy Suzuki":   {"color": "#34495E", "label": "⚡ [SUZUKI]\nSản xuất beam"},
        "Máy Sectional": {"color": "#34495E", "label": "⚙️ [MÁY SECTIONAL]\nSản xuất beam"},
        "Kho Beam Sizing": {"color": "#2C3E50", "label": "🧻 [KHO BEAM SIZING]\nLưu trục Sizing"},
        "Kho Sợi Weaving": {"color": "#2C3E50", "label": "🧵 [KHO SỢI WEAVING]\nKho sợi dệt"},
        "Kho Beam Weaving": {"color": "#2C3E50", "label": "🧻 [KHO BEAM WEAVING]\nKho trục dệt"},
        "Xưởng Dệt": {"color": "#2C3E50", "label": "🏭 [XƯỞNG DỆT]\nSản xuất vải"},
        "Kho Thành Phẩm": {"color": "#2C3E50", "label": "🛍️ [KHO THÀNH PHẨM]\nNhập kho vải"}
    }
    for name, props in d_nodes.items():
        graph.node(name, props["label"], fillcolor=props["color"], **node_attr)

    edge_attr = {'color': '#00FFAA', 'penwidth': '1.5', 'arrowsize': '1.0'}
    
    with graph.subgraph() as s:
        s.attr(rank='same')
        s.node('Kho Sợi Tổng')
        s.node('Kho Sợi Sizing')
    with graph.subgraph() as s:
        s.attr(rank='same')
        s.node('Xưởng Nhuộm')
        s.node('Máy Direct')
        s.node('Máy Sectional')
        s.node('Máy Winder')
        s.node('Máy Suzuki')
        s.node('Máy Winder')
        s.node('Máy Suzuki')
    with graph.subgraph() as s:
        s.attr(rank='same')
        s.node('Máy Hồ')
        s.node('Kho Sợi Weaving')
    with graph.subgraph() as s:
        s.attr(rank='same')
        s.node('Kho Beam Sizing')
        s.node('Kho Beam Weaving')
    with graph.subgraph() as s:
        s.attr(rank='same')
        s.node('Xưởng Dệt')
        s.node('Kho Thành Phẩm')

    edges = [
        ('Kho Sợi Tổng', 'Xưởng Nhuộm'), ('Kho Sợi Tổng', 'Kho Sợi Sizing'), ('Kho Sợi Tổng', 'Kho Sợi Weaving'),
        ('Xưởng Nhuộm', 'Kho Sợi Weaving'), ('Xưởng Nhuộm', 'Kho Sợi Sizing'),
        ('Kho Sợi Sizing', 'Máy Direct'), ('Máy Direct', 'Máy Hồ'), ('Máy Hồ', 'Kho Beam Sizing'),
        ('Kho Sợi Sizing', 'Máy Sectional'), ('Máy Sectional', 'Kho Beam Sizing'),
        ('Kho Beam Sizing', 'Kho Beam Weaving'),
        ('Kho Sợi Weaving', 'Xưởng Dệt'), ('Kho Beam Weaving', 'Xưởng Dệt'), ('Xưởng Dệt', 'Kho Thành Phẩm')
    ]
    for src, dst in edges:
        graph.edge(src, dst, **edge_attr)

    col1, col2, col3 = st.columns([1, 8, 1])
    with col2:
        st.graphviz_chart(graph, use_container_width=True)

    st.divider()

    # --- PHẦN 1: BÁO CÁO CHI TIẾT ---
    st.subheader("🔍 PHÂN TÍCH CHI TIẾT TỪNG CỤM")
    
    selected_node = st.selectbox("👉 Chọn Cụm/Kho trên sơ đồ để xem báo cáo:", list(d_nodes.keys()))
    
    df_cluster = pd.DataFrame() 
    current_view_date = None

    # ── Sizing Dashboard Renderer ─────────────────────────────────────────────
    def _render_sizing_dashboard(df_sz, node_name):
        import sqlite3 as _sl5

        # Machine icons mapping
        _MAC_ICON = {
            "SIZING BENNINGER": "🏭", "SIZING KARL MAYER": "⚙️",
            "SIZING HONGHWA": "🔧", "SECTIONAL": "🔩",
            "DIRECT": "🔗", "WINDER": "🎡", "SUZUKI": "⚡",
        }
        def _icon(name):
            n = str(name).upper()
            for k, v in _MAC_ICON.items():
                if k in n: return v
            return "🔨"

        if 'date' not in df_sz.columns or df_sz.empty:
            st.info("Chưa có dữ liệu"); return

        dates = sorted([str(d) for d in df_sz['date'].dropna().unique()], reverse=True)
        if not dates: st.info("Chưa có ngày"); return

        sel_date = st.selectbox("📅 Chọn ngày:", dates, key="sz_date_sel")
        df_day = df_sz[df_sz['date'] == sel_date].copy()

        # KPIs
        tot_mtr = df_day['sl_thuc_te_mtr'].sum() if 'sl_thuc_te_mtr' in df_day else 0
        tot_kg  = df_day['sl_kg'].sum()           if 'sl_kg'          in df_day else df_day.get('quantity_kg', pd.Series([0])).sum()
        avg_hs  = df_day['hieu_suat_pct'].mean()  if 'hieu_suat_pct'  in df_day else 0
        n_mac   = df_day['ten_may'].nunique()      if 'ten_may'        in df_day else 0

        c1,c2,c3,c4 = st.columns(4)
        with c1: st.metric("📏 Tổng Mét", f"{tot_mtr:,.0f} m")
        with c2: st.metric("⚖️ Tổng Kg",  f"{tot_kg:,.1f} kg")
        with c3: st.metric("📊 Hiệu Suất TB", f"{avg_hs:.1%}" if avg_hs else "—")
        with c4: st.metric("🏭 Số Máy", str(n_mac))

        st.markdown("---")
        st.markdown("#### 🏭 Tình Trạng Từng Máy")

        # Group by machine
        if 'ten_may' not in df_day.columns:
            st.dataframe(df_day, use_container_width=True, hide_index=True); return

        machines = df_day['ten_may'].dropna().unique()
        _mcols = st.columns(min(len(machines), 3))
        for _mi, _mac in enumerate(machines):
            df_m = df_day[df_day['ten_may'] == _mac]
            _ic  = _icon(_mac)
            _mtr = df_m['sl_thuc_te_mtr'].sum() if 'sl_thuc_te_mtr' in df_m else 0
            _kg  = df_m['sl_kg'].sum() if 'sl_kg' in df_m else 0
            _hs  = df_m['hieu_suat_pct'].mean() if 'hieu_suat_pct' in df_m else 0
            _n_ca = len(df_m)
            # Tốc độ thực tế
            _spd = df_m['toc_do_thuc_te'].mean() if 'toc_do_thuc_te' in df_m else 0
            # Bar màu theo hiệu suất
            _color = "#0CA678" if _hs > 0.6 else ("#F59F00" if _hs > 0.4 else "#FA5252")
            _pct   = int(_hs * 100) if _hs else 0
            # Time range
            _start = df_m['gio_bat_dau'].min() if 'gio_bat_dau' in df_m else ""
            _end   = df_m['gio_ket_thuc'].max() if 'gio_ket_thuc' in df_m else ""
            _time  = f"{_start} → {_end}" if _start and _end else "—"

            with _mcols[_mi % 3]:
                st.markdown(f"""
<div style="background:#fff;border:1px solid #E4E7EE;border-radius:10px;
            padding:16px;margin-bottom:12px;box-shadow:0 1px 4px rgba(0,0,0,0.06)">
  <div style="display:flex;align-items:center;gap:8px;margin-bottom:10px">
    <span style="font-size:22px">{_ic}</span>
    <div>
      <div style="font-weight:700;font-size:14px;color:#1A1B2E">{_mac}</div>
      <div style="font-size:11px;color:#8B92A5">{_n_ca} ca | {_time}</div>
    </div>
  </div>
  <div style="display:grid;grid-template-columns:1fr 1fr;gap:6px;margin-bottom:10px">
    <div style="background:#F0F2F5;border-radius:6px;padding:8px;text-align:center">
      <div style="font-size:11px;color:#8B92A5;font-weight:600;text-transform:uppercase;letter-spacing:0.5px">Mét</div>
      <div style="font-size:16px;font-weight:700;color:#1A1B2E">{_mtr:,.0f}</div>
    </div>
    <div style="background:#F0F2F5;border-radius:6px;padding:8px;text-align:center">
      <div style="font-size:11px;color:#8B92A5;font-weight:600;text-transform:uppercase;letter-spacing:0.5px">Kg</div>
      <div style="font-size:16px;font-weight:700;color:#1A1B2E">{_kg:,.1f}</div>
    </div>
  </div>
  <div style="margin-bottom:4px;display:flex;justify-content:space-between">
    <span style="font-size:11px;color:#8B92A5;font-weight:600">HIỆU SUẤT</span>
    <span style="font-size:12px;font-weight:700;color:{_color}">{_pct}%</span>
  </div>
  <div style="background:#E4E7EE;border-radius:100px;height:5px;overflow:hidden">
    <div style="background:{_color};width:{_pct}%;height:100%;border-radius:100px;
                transition:width 0.4s ease"></div>
  </div>
  {"<div style='margin-top:6px;font-size:11px;color:#8B92A5'>Tốc độ: " + str(round(_spd)) + " m/min</div>" if _spd else ""}
</div>
""", unsafe_allow_html=True)

        # Details expander
        with st.expander("📋 Chi tiết từng ca"):
            _show_cols = [c for c in ['ten_may','ca','loai_soi','pd_yd','toc_do_thuc_te',
                                       'sl_thuc_te_mtr','sl_kg','hieu_suat_pct',
                                       'gio_bat_dau','gio_ket_thuc','thoi_gian_phut','ghi_chu']
                          if c in df_day.columns]
            st.dataframe(df_day[_show_cols].rename(columns={
                'ten_may':'Máy','ca':'Ca','loai_soi':'Loại sợi',
                'pd_yd':'PD/YD','toc_do_thuc_te':'Tốc độ',
                'sl_thuc_te_mtr':'Mét','sl_kg':'Kg',
                'hieu_suat_pct':'HS%','gio_bat_dau':'Bắt đầu',
                'gio_ket_thuc':'Kết thúc','thoi_gian_phut':'Phút','ghi_chu':'Ghi chú'
            }).sort_values('Máy'), use_container_width=True, hide_index=True)
    # ─────────────────────────────────────────────────────────────────────────

    with st.container(border=True):
        st.markdown(f"### 📈 Báo cáo: {selected_node}")
        
        try:
            # ── Route: Sizing nodes → Sizing_Log, others → Inventory_Log ──
            _SIZING_MAP = {
                "Máy Hồ": "MÁY HỒ", "Máy Sectional": "MÁY SEC",
                "Máy Direct": "MÁY QS", "Máy Winder": "WINDER", "Máy Suzuki": "SUZUKI",
            }
            _is_sizing = selected_node in _SIZING_MAP

            if _is_sizing:
                _mtype = _SIZING_MAP[selected_node]
                df_cluster_raw = execute_query(f"""
                    SELECT date, machine_type AS cluster_name,
                           ten_may, ca, nguoi_chay, loai_soi, pd_yd,
                           toc_do_muc_tieu, toc_do_thuc_te,
                           sl_muc_tieu_mtr, sl_thuc_te_mtr, sl_kg,
                           hieu_suat_pct, hs_toc_do_pct,
                           gio_bat_dau, gio_ket_thuc, thoi_gian_phut,
                           dut_soi, sua_chua
                    FROM Sizing_Log WHERE machine_type = '{_mtype}'
                """)
            else:
                df_cluster_raw = execute_query(f"SELECT * FROM Inventory_Log WHERE cluster_name = '{selected_node}'")

            
            if df_cluster_raw.empty:
                st.info(f"Chưa có dữ liệu nào được tải lên cho cụm [{selected_node}]. Hãy cấu hình Auto-Sync để nạp file.")
            else:
                if _is_sizing:
                    _render_sizing_dashboard(df_cluster_raw, selected_node)
                    df_cluster = df_cluster_raw
                else:
                    df_cluster = get_clean_cluster_data(df_cluster_raw, selected_node)

                if not _is_sizing and selected_node == "Xưởng Dệt":
                    mac_col = find_best_col(df_cluster, ['ten_may', 'số máy', 'máy', 'mac'])
                    loc_col = find_best_col(df_cluster, ['sub_location', 'khu vực', 'xưởng', 'sheet', 'vị trí'])
                    y_col = find_best_col(df_cluster, ['quantity_yard', 'yard', 'mét', 'met'])
                    kg_col = find_best_col(df_cluster, ['quantity_kg', 'tổng', 'total', 'kg', 'khối lượng'])
                    order_col = find_best_col(df_cluster, ['order_id', 'tên hàng', 'sản phẩm', 'mã', 'item_id'])

                    if not mac_col:
                        st.error("🚨 HỆ THỐNG KHÔNG TÌM THẤY CỘT SỐ MÁY. Sếp hãy kiểm tra lại cấu hình Mapping, hãy đảm bảo cột chứa Số Máy được map hoặc có tên chứa chữ 'máy'.")
                    else:
                        if 'date' not in df_cluster.columns: df_cluster['date'] = pd.Timestamp.now().strftime('%Y-%m-%d')
                        # CHỐNG LỖI KHI TÌM UNIQUE TRÊN KIỂU DỮ LIỆU TRỘN LẪN
                        available_dates = sorted([str(d) for d in df_cluster['date'].replace("", pd.NA).dropna().unique()], reverse=True)
                        
                        if not available_dates: 
                            st.info("Chưa có ngày sản xuất nào được ghi nhận hợp lệ.")
                        else:
                            selected_date = st.selectbox("Chọn ngày xem báo cáo Xưởng Dệt:", available_dates)
                            current_view_date = selected_date
                            df_day = df_cluster[df_cluster['date'] == selected_date].copy()

                            if y_col and y_col in df_day.columns: 
                                df_day[y_col] = pd.to_numeric(df_day[y_col], errors='coerce').fillna(0)
                            if kg_col and kg_col in df_day.columns: 
                                df_day[kg_col] = pd.to_numeric(df_day[kg_col], errors='coerce').fillna(0)
                            
                            total_yard = df_day[y_col].sum() if y_col and y_col in df_day.columns and not df_day.empty else 0
                            total_kg = df_day[kg_col].sum() if kg_col and kg_col in df_day.columns and not df_day.empty else 0
                            
                            active_machines_df = df_day[~df_day[order_col].astype(str).str.lower().isin(["nan", "none", "", "0", "chưa xác định", "pd.na"])] if order_col and order_col in df_day.columns and not df_day.empty else pd.DataFrame()
                            active_machines = active_machines_df[mac_col].nunique() if not active_machines_df.empty else 0
                            
                            col1_1, col1_2, col1_3 = st.columns(3)
                            with col1_1: st.metric("Tổng Sản Lượng (Yard)", f"{total_yard:,.0f} Y")
                            with col1_2: st.metric("Tổng Sản Lượng (Kg)", f"{total_kg:,.1f} Kg")
                            with col1_3: st.metric("Số Máy Đang Chạy", f"{active_machines} máy")

                            st.markdown("---")
                            st.markdown("### 🏭 SƠ ĐỒ MẶT BẰNG XƯỞNG DỆT (DIGITAL TWIN)")
                            
                            def assign_zone(loc_str, mac_str):
                                loc_str = str(loc_str).lower().strip()
                                mac_str = str(mac_str).lower().strip()
                                if loc_str in ['1', '1.0', 'wea 1', 'wea1', 'weaving 1', 'w1'] or 'w1' in mac_str or 'wea 1' in loc_str: return 1
                                if loc_str in ['2', '2.0', 'wea 2', 'wea2', 'weaving 2', 'w2'] or 'w2' in mac_str or 'wea 2' in loc_str: return 2
                                if loc_str in ['3', '3.0', 'wea 3', 'wea3', 'weaving 3', 'w3'] or 'w3' in mac_str or 'wea 3' in loc_str: return 3
                                try:
                                    val = int(float(loc_str))
                                    if val in [1, 2, 3]: return val
                                except: pass
                                return 0

                            df_day['zone_id'] = df_day.apply(lambda row: assign_zone(row.get(loc_col, ''), row.get(mac_col, '')), axis=1)

                            unknown_mask = df_day['zone_id'] == 0
                            if not df_day.empty and unknown_mask.any():
                                z_counts = df_day[~unknown_mask].groupby('zone_id')[mac_col].nunique().to_dict()
                                for idx in df_day[unknown_mask].index:
                                    if z_counts.get(1, 0) < 48: 
                                        df_day.at[idx, 'zone_id'] = 1
                                        z_counts[1] = z_counts.get(1, 0) + 1
                                    elif z_counts.get(2, 0) < 56: 
                                        df_day.at[idx, 'zone_id'] = 2
                                        z_counts[2] = z_counts.get(2, 0) + 1
                                    else: 
                                        df_day.at[idx, 'zone_id'] = 3
                                        z_counts[3] = z_counts.get(3, 0) + 1

                            tab_w1, tab_w2, tab_w3 = st.tabs(["🏭 Weaving 1 (48 Máy)", "🏭 Weaving 2 (56 Máy)", "🏭 Weaving 3 (24 Máy)"])
                            
                            def render_zone_dashboard(zone_id, capacity, title, color_hex):
                                df_zone = df_day[df_day['zone_id'] == zone_id].copy() if 'zone_id' in df_day.columns and not df_day.empty else pd.DataFrame()
                                
                                total_y_zone = df_zone[y_col].sum() if y_col and y_col in df_zone.columns and not df_zone.empty else 0
                                total_kg_zone = df_zone[kg_col].sum() if kg_col and kg_col in df_zone.columns and not df_zone.empty else 0
                                
                                active_macs_df_zone = df_zone[~df_zone[order_col].astype(str).str.lower().isin(["nan", "none", "", "0", "chưa xác định", "pd.na"])] if order_col and order_col in df_zone.columns and not df_zone.empty else pd.DataFrame()
                                active_macs_zone = active_macs_df_zone[mac_col].nunique() if mac_col and mac_col in active_macs_df_zone.columns and not active_macs_df_zone.empty else 0
                                
                                c1, c2, c3 = st.columns(3)
                                with c1: st.metric(f"Sản Lượng {title} (Yard)", f"{total_y_zone:,.0f} Y")
                                with c2: st.metric(f"Sản Lượng {title} (Kg)", f"{total_kg_zone:,.1f} Kg")
                                with c3: st.metric("Số Máy Đang Chạy", f"{active_macs_zone} máy")
                                
                                st.markdown("---")
                                
                                zone_mac_list = []
                                z_grouped = pd.DataFrame()
                                
                                if not df_zone.empty:
                                    agg_funcs = {}
                                    if y_col and y_col in df_zone.columns: agg_funcs[y_col] = 'sum'
                                    if kg_col and kg_col in df_zone.columns: agg_funcs[kg_col] = 'sum'
                                    if order_col and order_col in df_zone.columns: agg_funcs[order_col] = 'first'
                                    
                                    group_keys = [mac_col, loc_col] if loc_col and loc_col in df_zone.columns else [mac_col]
                                    
                                    if agg_funcs: z_grouped = df_zone.groupby(group_keys).agg(agg_funcs).reset_index()
                                    else: z_grouped = df_zone[group_keys].drop_duplicates()
                                    
                                    # CHỐNG LỖI "min() arg is an empty sequence" KHI GROUPBY RỖNG
                                    if not z_grouped.empty:
                                        z_grouped['sort_key'] = pd.to_numeric(z_grouped[mac_col].astype(str).str.extract(r'(\d+)')[0], errors='coerce').fillna(999)
                                        z_grouped = z_grouped.sort_values('sort_key')
                                    
                                        for _, row in z_grouped.iterrows():
                                            raw_mac_name = str(row[mac_col]).strip()
                                            if raw_mac_name.endswith('.0'): raw_mac_name = raw_mac_name[:-2]
                                            if raw_mac_name.isdigit(): raw_mac_name = f"Máy {raw_mac_name}"
                                            
                                            zone_mac_list.append({
                                                "mac_name": raw_mac_name,
                                                "yard": row.get(y_col, 0) if y_col else 0,
                                                "kg": row.get(kg_col, 0) if kg_col else 0,
                                                "order": str(row.get(order_col, "N/A")).strip() if order_col else "N/A"
                                            })
                                        
                                running_macs = [m for m in zone_mac_list if str(m['order']).lower() not in ["nan", "none", "", "0", "chưa xác định", "pd.na"]]
                                used_slots = len(running_macs)
                                
                                html = f"<h5 style='margin-top: 5px; margin-bottom: 15px;'>Sơ đồ {title} (Đang chạy: {used_slots} / Tổng máy: {capacity})</h5>"
                                html += "<div style='display: grid; grid-template-columns: repeat(auto-fill, minmax(65px, 1fr)); gap: 6px; margin-bottom: 25px;'>"
                                for i in range(capacity):
                                    if i < len(zone_mac_list):
                                        mac = zone_mac_list[i]
                                        m_order, m_yard, m_kg, m_name = mac['order'], mac['yard'], mac['kg'], mac['mac_name']
                                        if str(m_order).lower() in ["nan", "none", "", "0", "chưa xác định", "pd.na"]:
                                            html += (
                                                f"<div title='{m_name} đang dừng' style='"
                                                f"background-color: #E2E8F0; color: #718096; "
                                                f"padding: 6px 3px; border-radius: 6px; text-align: center; "
                                                f"font-size: 10px; border: 1px dashed #A0AEC0; "
                                                f"min-height: 70px; display:flex; flex-direction:column; "
                                                f"align-items:center; justify-content:center;'>"
                                                f"<b style='font-size:11px; color:#4A5568'>{m_name}</b>"
                                                f"<span style='font-size:9px'>Dừng</span>"
                                                f"</div>"
                                            )
                                        else:
                                            # Rút gọn tên hàng cho vừa ô
                                            _order_short = str(m_order)[:18] + "…" if len(str(m_order)) > 18 else str(m_order)
                                            _kg_str = f"{m_kg:,.0f} Kg" if m_kg else ""
                                            _tooltip = f"Mã hàng: {m_order}\nSản lượng: {m_kg:,.1f} Kg | Yard: {m_yard:,.0f}"
                                            html += (
                                                f"<div title='{_tooltip}' style='"
                                                f"background-color: {color_hex}; color: white; "
                                                f"padding: 6px 3px; border-radius: 6px; text-align: center; "
                                                f"font-size: 10px; cursor: pointer; "
                                                f"border: 1px solid rgba(255,255,255,0.4); "
                                                f"box-shadow: 0 2px 4px rgba(0,0,0,0.2); "
                                                f"line-height: 1.3; min-height: 70px; "
                                                f"display:flex; flex-direction:column; align-items:center; justify-content:center;'>"
                                                f"<b style='font-size:11px'>{m_name}</b>"
                                                f"<span style='font-size:9px; opacity:0.9; word-break:break-word; margin-top:2px'>{_order_short}</span>"
                                                f"<span style='font-size:10px; font-weight:bold; margin-top:2px'>{_kg_str}</span>"
                                                f"</div>"
                                            )
                                    else:
                                        html += (
                                                "<div title='Không có dữ liệu máy' style='"
                                                "background-color: #F7FAFC; color: #CBD5E0; "
                                                "padding: 6px 3px; border-radius: 6px; text-align: center; "
                                                "font-size: 9px; border: 1px dashed #E2E8F0; "
                                                "min-height: 70px; display:flex; align-items:center; justify-content:center;'>"
                                                "Trống</div>"
                                            )
                                html += "</div>"
                                st.markdown(html, unsafe_allow_html=True)
                                
                                with st.expander(f"Xem bảng số liệu chi tiết của riêng {title}"):
                                    # CHỐNG LỖI st.dataframe RỖNG
                                    if not df_zone.empty:
                                        st.dataframe(df_zone.drop(columns=['id', 'cluster_name', 'type', 'zone_id'], errors='ignore'), use_container_width=True)
                                    else:
                                        st.info(f"Không có chi tiết nào cho khu vực {title}.")

                            with tab_w1: render_zone_dashboard(1, 48, "Weaving 1", "#D69E2E")
                            with tab_w2: render_zone_dashboard(2, 56, "Weaving 2", "#3182CE")
                            with tab_w3: render_zone_dashboard(3, 24, "Weaving 3", "#00B517")

                elif "Kho" in selected_node:
                    qty_col = find_best_col(df_cluster, ['quantity', 'số lượng', 'mét', 'met', 'kg', 'total', 'tổng'])
                    item_col = find_best_col(df_cluster, ['item_id', 'tên hàng', 'mã', 'sản phẩm', 'order_id'])
                    loc_col = find_best_col(df_cluster, ['sub_location', 'khu vực', 'vị trí', 'kho'])

                    if not qty_col or not item_col:
                        st.error("🚨 HỆ THỐNG KHÔNG TÌM THẤY CỘT 'MÃ HÀNG' HOẶC 'SỐ LƯỢNG'. Hãy kiểm tra lại Mapping.")
                    else:
                        df_cluster[qty_col] = pd.to_numeric(df_cluster[qty_col], errors='coerce').fillna(0)
                        
                        col1_1, col1_2, col1_3 = st.columns(3)
                        if not loc_col:
                            loc_col = 'sub_location_temp'
                            df_cluster[loc_col] = "Khu vực dưới đất"
                        else:
                            df_cluster[loc_col] = df_cluster[loc_col].replace("", "Khu vực dưới đất").fillna("Khu vực dưới đất")
                        
                        valid_items = df_cluster[df_cluster[item_col].astype(str) != "Chưa xác định"]
                        df_nhap = valid_items[valid_items['type'].isin(['NHAP', 'TON_DAU'])].groupby([item_col, loc_col])[qty_col].sum()
                        df_xuat = valid_items[valid_items['type'] == 'XUAT'].groupby([item_col, loc_col])[qty_col].sum()
                        df_stock = df_nhap.sub(df_xuat, fill_value=0).reset_index(name='quantity')
                        df_stock = df_stock[df_stock['quantity'] > 0] 
                        
                        ton_kho_tong = df_stock['quantity'].sum()
                        so_luong_ma_sp = df_stock[item_col].nunique()
                        
                        with col1_1: st.metric("Tổng Tồn Kho Hiện Tại", f"{ton_kho_tong:,.0f}")
                        with col1_2: st.metric("Số Mã Sản Phẩm Đang Lưu", f"{so_luong_ma_sp}")
                        with col1_3: st.metric("Cảnh báo Sắp Hết Hạn", "0", delta="- Bình thường", delta_color="normal")

                        if selected_node == "Kho Beam Weaving":
                            # ✅ PATH INPUT + 1-CLICK SYNC cho Kho Beam Weaving
                            if "kbw_file_path" not in st.session_state:
                                st.session_state.kbw_file_path = st.session_state.get("beam_file_path", "")
                            if "kbw_last_synced" not in st.session_state:
                                st.session_state.kbw_last_synced = ""

                            with st.container(border=True):
                                st.caption("🔗 **Tự động cập nhật Kho Beam Weaving từ file BEAM_DAT**")
                                _c1, _c2, _c3 = st.columns([4, 1.5, 1.5])
                                with _c1:
                                    _kbw_path = st.text_input(
                                        "Đường dẫn file BEAM_DAT:",
                                        value=st.session_state.kbw_file_path,
                                        key="kbw_path_input",
                                        placeholder="VD: Z:\VINA_folder\BEAM_DAT-THONG_TIN_BEAM_2025.xlsx",
                                        label_visibility="collapsed",
                                    )
                                    if _kbw_path != st.session_state.kbw_file_path:
                                        st.session_state.kbw_file_path = _kbw_path
                                        st.session_state.kbw_last_synced = ""

                                with _c2:
                                    _btn_kbw = st.button(
                                        "🔄 Cập nhật kho", key="btn_kbw_sync",
                                        use_container_width=True,
                                        disabled=(not _kbw_path or not os.path.isfile(_kbw_path))
                                    )
                                with _c3:
                                    if st.session_state.kbw_last_synced:
                                        st.caption(f"✅ Đã sync: {st.session_state.kbw_last_synced}")
                                    elif _kbw_path and os.path.isfile(_kbw_path):
                                        st.caption("⏳ Chưa sync — bấm Cập nhật")
                                    elif _kbw_path:
                                        st.caption("⚠️ Không tìm thấy file")

                                if _btn_kbw and _kbw_path and os.path.isfile(_kbw_path):
                                    with st.spinner("Đang đọc & tính diff vs hôm qua..."):
                                        try:
                                            from kho_beam_pipeline import import_snapshot as _kbw_snap
                                            _df_snap = pd.read_excel(_kbw_path, header=1)
                                            _snap_date = pd.Timestamp.today().strftime("%Y-%m-%d")
                                            _res_kbw = _kbw_snap(_df_snap, _snap_date, os.path.basename(_kbw_path))
                                            import time as _tm_kbw
                                            st.session_state.kbw_last_synced = _tm_kbw.strftime("%d/%m %H:%M")
                                            if _res_kbw.get("error"):
                                                st.error(_res_kbw["error"])
                                            else:
                                                st.success(
                                                    f"✅ Snapshot {_snap_date}: "
                                                    f"**{_res_kbw['snapshot']} beam** trong kho | "
                                                    f"🟢 Nhập: {_res_kbw['nhap']} | "
                                                    f"🔴 Xuất: {_res_kbw['xuat']}"
                                                    + (f" (so với {_res_kbw['prev_date']})" if _res_kbw.get('prev_date') else " (snapshot đầu tiên)")
                                                )
                                            st.rerun()
                                        except Exception as _ekbw:
                                            st.error(f"Lỗi: {_ekbw}")

                                # Upload snapshot thủ công (từ file ERP export)
                                st.caption("📤 Hoặc tải file snapshot thủ công (export từ ERP):")
                                _cc1, _cc2 = st.columns([3, 2])
                                with _cc1:
                                    _kbw_upload = st.file_uploader(
                                        "File ERP snapshot:", type=["xlsx","xls","csv"],
                                        key="kbw_upload_snap", label_visibility="collapsed"
                                    )
                                with _cc2:
                                    _kbw_snap_date = st.date_input(
                                        "Ngày snapshot:", value=pd.Timestamp.today(),
                                        key="kbw_snap_date_input", label_visibility="collapsed"
                                    )
                                if _kbw_upload and st.button("📥 Nhập snapshot này", key="btn_kbw_upload_snap"):
                                    with st.spinner("Đang xử lý diff..."):
                                        try:
                                            from kho_beam_pipeline import import_snapshot as _kbw_snap2
                                            if _kbw_upload.name.endswith(".csv"):
                                                _df_up = pd.read_csv(_kbw_upload)
                                            else:
                                                _df_up = pd.read_excel(_kbw_upload, header=1)
                                            _date_str = str(_kbw_snap_date)
                                            _res2 = _kbw_snap2(_df_up, _date_str, _kbw_upload.name)
                                            if _res2.get("error"):
                                                st.error(_res2["error"])
                                            else:
                                                st.success(
                                                    f"✅ {_date_str}: {_res2['snapshot']} beam | "
                                                    f"Nhập {_res2['nhap']} | Xuất {_res2['xuat']}"
                                                )
                                                if _res2.get("nhap_list"):
                                                    st.caption("Beam nhập: " + ", ".join(_res2["nhap_list"][:10]))
                                                if _res2.get("xuat_list"):
                                                    st.caption("Beam xuất: " + ", ".join(_res2["xuat_list"][:10]))
                                            st.rerun()
                                        except Exception as _eu:
                                            st.error(f"Lỗi: {_eu}")

                            st.markdown("---")
                            st.markdown("### 🗄️ SƠ ĐỒ LƯU TRỮ BEAM THỰC TẾ (SỐ HÓA)")
                            
                            def extract_location(loc_str):
                                loc_str = str(loc_str).strip()
                                match = re.search(r'(60|95)\D+(\d+)', loc_str)
                                if match: return int(match.group(1)), int(match.group(2))
                                return None, None

                            rack_60_slots, rack_95_slots, floor_beams = [None] * 60, [None] * 95, []
                            for _, row in df_stock.iterrows():
                                beam_data = row.to_dict()
                                rack_type, slot_idx = extract_location(beam_data[loc_col])
                                if rack_type == 60 and 1 <= slot_idx <= 60: rack_60_slots[slot_idx - 1] = beam_data
                                elif rack_type == 95 and 1 <= slot_idx <= 95: rack_95_slots[slot_idx - 1] = beam_data
                                else: floor_beams.append(beam_data) 

                            def draw_fixed_rack(title, capacity, slot_array, color_hex):
                                used_slots = sum(1 for x in slot_array if x is not None)
                                html = f"<h5 style='margin-top: 5px; margin-bottom: 15px;'>{title} (Đang dùng: {used_slots} / Sức chứa: {capacity})</h5>"
                                html += "<div style='display: grid; grid-template-columns: repeat(auto-fill, minmax(60px, 1fr)); gap: 6px; margin-bottom: 25px;'>"
                                for i in range(capacity):
                                    beam = slot_array[i]
                                    if beam is not None:
                                        short_id = str(beam[item_col])[-5:] if len(str(beam[item_col])) > 5 else str(beam[item_col])
                                        b_id, b_qty = beam[item_col], beam['quantity']
                                        html += f"<div title='Mã Beam: {b_id} | Tọa độ: Ô số {i+1} | SL: {b_qty:,.0f}' style='background-color: {color_hex}; color: white; padding: 12px 2px; border-radius: 4px; text-align: center; font-size: 11px; cursor: pointer; border: 1px solid rgba(255,255,255,0.4); font-weight: bold;'>{short_id}</div>"
                                    else:
                                        html += f"<div title='Vị trí Ô {i+1} trống' style='background-color: #1A202C; color: #4A5568; padding: 12px 2px; border-radius: 4px; text-align: center; font-size: 10px; border: 1px dashed #4A5568;'>Trống {i+1}</div>"
                                html += "</div>"
                                return html
                            
                            def draw_floor(beams_list):
                                html = f"<h5 style='margin-top: 5px; margin-bottom: 15px;'>Khu Vực Dưới Đất (Đang để: {len(beams_list)} cuộn)</h5>"
                                if len(beams_list) == 0: return html + "<p style='color:#718096; font-style:italic;'>Kho trống.</p>"
                                html += "<div style='display: grid; grid-template-columns: repeat(auto-fill, minmax(60px, 1fr)); gap: 6px; margin-bottom: 25px;'>"
                                for beam in beams_list:
                                    short_id = str(beam[item_col])[-5:] if len(str(beam[item_col])) > 5 else str(beam[item_col])
                                    b_id, b_loc, b_qty = beam[item_col], beam[loc_col], beam['quantity']
                                    html += f"<div title='Mã: {b_id} | Vị trí: {b_loc} | SL: {b_qty:,.0f}' style='background-color: #D69E2E; color: white; padding: 12px 2px; border-radius: 4px; text-align: center; font-size: 11px; cursor: pointer; border: 1px solid rgba(255,255,255,0.4); font-weight: bold;'>{short_id}</div>"
                                html += "</div>"
                                return html

                            tab_all, tab_60, tab_95, tab_floor = st.tabs(["👁️ Tất cả Kho", "🟢 Giá 60 chỗ", "🔵 Giá 95 chỗ", "🟡 Khu dưới đất"])
                            with tab_all:
                                st.markdown(draw_fixed_rack("1. Giá Treo Beam (Khu Vực 60 Chỗ)", 60, rack_60_slots, "#00B517"), unsafe_allow_html=True)
                                st.markdown(draw_fixed_rack("2. Giá Treo Beam (Khu Vực 95 Chỗ)", 95, rack_95_slots, "#3182CE"), unsafe_allow_html=True)
                                st.markdown(draw_floor(floor_beams), unsafe_allow_html=True)
                            with tab_60: st.markdown(draw_fixed_rack("Giá Treo Beam 60 Chỗ", 60, rack_60_slots, "#00B517"), unsafe_allow_html=True)
                            with tab_95: st.markdown(draw_fixed_rack("Giá Treo Beam 95 Chỗ", 95, rack_95_slots, "#3182CE"), unsafe_allow_html=True)
                            with tab_floor: st.markdown(draw_floor(floor_beams), unsafe_allow_html=True)
                        else:
                            st.markdown("**📊 Biểu đồ Tồn theo Sản Phẩm (Top 10)**")
                            df_stock_total = df_stock.groupby(item_col)['quantity'].sum().reset_index()
                            # CHỐNG LỖI BIỂU ĐỒ RỖNG
                            if not df_stock_total.empty: st.bar_chart(df_stock_total.sort_values(by='quantity', ascending=False).head(10).set_index(item_col)['quantity'])

                with st.expander("Xem bảng số liệu chi tiết (Gốc của Cụm này)"):
                    df_clean = df_cluster.copy()
                    
                    try:
                        db_files = glob.glob("*.db")
                        if db_files and "Kho" in selected_node:
                            with sqlite3.connect(db_files[0]) as conn:
                                df_yarn = pd.read_sql_query("SELECT * FROM Yarn_Dictionary", conn)
                                coeff_dict = {str(k).strip().upper(): float(v) for k, v in zip(df_yarn['yarn_type'], df_yarn['coefficient'])} if not df_yarn.empty else {}
                                
                                qty_col = find_best_col(df_clean, ['quantity', 'mét', 'met', 'số lượng'])
                                yarn_col = find_best_col(df_clean, ['sợi', 'yarn'])
                                beam_size_col = find_best_col(df_clean, ['tổng', 'size'])
                                
                                if qty_col and yarn_col and beam_size_col and coeff_dict:
                                    kg_list = []
                                    for _, row in df_clean.iterrows():
                                        kg = ""
                                        try:
                                            yarn_val = str(row[yarn_col]).upper()
                                            core_match = re.search(r'(\d+[S]/\d+|\d+NTW)', yarn_val)
                                            core_yarn = core_match.group(1) if core_match else yarn_val
                                            coeff = coeff_dict.get(core_yarn)
                                            if not coeff: 
                                                for k, v in coeff_dict.items():
                                                    if k in core_yarn or core_yarn in k: coeff = v; break
                                            if coeff:
                                                m_str = str(row[qty_col]).replace(',', '').strip()
                                                ts_str = str(row[beam_size_col]).replace(',', '').strip()
                                                if m_str and ts_str and m_str != "nan" and ts_str != "nan":
                                                    kg = math.ceil(((float(m_str) / 0.9144) * float(ts_str) * coeff) + 10) / 1000
                                        except: pass
                                        kg_list.append(kg)
                                    df_clean['Kg_TonKho'] = kg_list
                    except Exception: pass
                    
                    # CHỐNG LỖI BẢNG GỐC RỖNG
                    if not df_clean.empty:
                        st.dataframe(df_clean.sort_values('date', ascending=False) if 'date' in df_clean.columns else df_clean, use_container_width=True)
                    else:
                        st.info("Bảng dữ liệu gốc trống.")

        except Exception as e:
            st.error(f"Lỗi tải dữ liệu chi tiết: {e}")

    st.markdown("<br>", unsafe_allow_html=True)

    # --- PHẦN 2: TRỢ LÝ AI ---
    st.subheader("🤖 TRỢ LÝ AI (Động cơ Lật Trang Tự Động)")
    # Model name warning (non-blocking — beam queries work without model)
    if not model_name:
        col_warn1, col_warn2 = st.columns([3,1])
        with col_warn1:
            st.warning("⚠️ Model AI trống — nhập tên model Ollama để dùng tính năng phân tích. Beam/sản lượng vẫn hoạt động.")
        with col_warn2:
            if st.button("🔧 Dùng qwen2.5:3b", key="btn_default_model", use_container_width=True):
                st.session_state.model_name = "qwen2.5:3b"
                st.rerun()

    if True:  # always render chat (model_name optional for beam queries)

        col_h1, col_h2 = st.columns([9,1])
        with col_h2:
            if st.button("🗑️ Xóa lịch sử", key="clr_mem", use_container_width=True):
                st.session_state.messages = []
                _save_memory([])
                st.rerun()
        # ── GỢI Ý CÂU HỎI ─────────────────────────────────────────────
        if "suggested_q" not in st.session_state: st.session_state.suggested_q = ""
        _QUESTIONS = {
            "🧵 Xưởng Dệt": [
                # 📦 Sản lượng — ngày cụ thể
                "Máy nào xưởng 1 hôm nay sản lượng cao nhất?",
                "Máy 27 xưởng 2 ngày 3/3 chạy được bao nhiêu?",
                "Tổng sản lượng xưởng 3 hôm nay bao nhiêu kg?",
                "Ngày nào xưởng 2 sản lượng cao nhất tháng 4?",
                "Xưởng 1 ngày 15/4 sản xuất được bao nhiêu kg?",
                # 📦 Sản lượng — tháng / kỳ
                "Sản lượng tổng xưởng dệt tháng 3/2026?",
                "Tổng sản lượng toàn nhà máy tháng 4/2026?",
                "Tháng nào xưởng 1 có sản lượng cao nhất 2026?",
                "Xưởng 2 tháng 4/2026 sản lượng bao nhiêu?",
                "Sản lượng xưởng 3 tháng này so tháng trước?",
                # ⚡ Hiệu suất
                "Hiệu suất trung bình xưởng 2 tháng này?",
                "Máy nào PD xưởng 1 tháng 3 sản lượng thấp nhất?",
                "Máy nào YD xưởng 2 hiệu suất cao nhất tháng 4?",
                "Hiệu suất YD so với PD xưởng 1 tháng 4?",
                "Máy nào xưởng 1 có hiệu suất dưới 50% tháng 3?",
                "Hiệu suất ca A xưởng 2 tháng 4 bao nhiêu?",
                # 📈 Xu hướng
                "Xu hướng sản lượng 10 ngày gần nhất xưởng 3",
                "Xu hướng hiệu suất 2 tuần qua xưởng 1?",
                "Sản lượng xưởng 2 7 ngày qua có tăng không?",
                "Máy nào xưởng 1 tháng này tiến bộ nhiều nhất?",
                # 🏷️ Mã hàng
                "Mã hàng nào đang chạy nhiều nhất xưởng 1?",
                "Mã hàng CS 32 tuần trước sản lượng bao nhiêu?",
                "Mã hàng SW LIGHT MUJI 40 tháng 4 chạy bao nhiêu kg?",
                "Top 5 mã hàng sản lượng cao nhất xưởng 2 tháng 3?",
                "Mã hàng nào ít nhất xưởng 3 tháng 4?",
                "Mã SW NEW COLOR MUJI 40 đang chạy ở máy nào?",
                # 🔄 So sánh xưởng
                "So sánh hiệu suất xưởng 1 và xưởng 2 tháng 3",
                "Xưởng nào sản lượng cao nhất tháng 4/2026?",
                "So sánh sản lượng tháng 3, 4, 5/2026 xưởng 2?",
            ],
            "🔩 Kho & Beam": [
                # 📏 Beam còn lại trên máy
                "Máy 1 xưởng 1 beam còn bao nhiêu mét?",
                "Beam máy 12 xưởng 1 còn bao nhiêu kg?",
                "Beam máy 5 xưởng 3 còn lại bao nhiêu phần trăm?",
                "Beam trên weaving 2 còn dưới 20% là những máy nào?",
                "Beam nào sắp hết trên xưởng 2?",
                "Beam máy 27 xưởng 2 dùng loại sợi gì?",
                "Tổng mét beam còn lại toàn xưởng 1?",
                # 🏷️ Thông tin beam lên máy
                "Beam nào lên máy gần nhất trên weaving 3?",
                "Beam trên máy 5 xưởng 3 lên ngày mấy?",
                "Beam nào lên máy trong tuần này toàn xưởng?",
                "Máy nào xưởng 2 chưa lên beam mới tháng này?",
                "Mã beam 26-2808 đang ở máy nào?",
                # 📦 Kho beam weaving (giá treo)
                "Tổng kg trong kho beam hiện tại?",
                "Kho beam weaving còn bao nhiêu beam?",
                "Giá 60 còn bao nhiêu ô trống?",
                "Giá 95 đang chứa bao nhiêu beam PD?",
                "Beam nào trong kho chưa lên máy xưởng 1?",
                # 🔗 Liên kết kho → xưởng
                "Beam lên máy tuần này từ kho là những beam nào?",
                "Tháng 4 nhập về kho beam bao nhiêu beam?",
                "Tháng 4 xuất khỏi kho beam bao nhiêu beam?",
            ],
            "⚙️ Sizing / Sectional": [
                # ⚡ Hiệu suất máy hồ
                "Hiệu suất máy hồ Bng tháng 3 là bao nhiêu?",
                "Hiệu suất trung bình máy sectional tháng 4/2026?",
                "Máy nào sizing có hiệu suất thấp nhất tháng này?",
                "Máy hồ Bng tháng 4 hiệu suất có cải thiện không?",
                "So sánh hiệu suất máy hồ tháng 3 và tháng 4?",
                # 📦 Sản lượng sizing
                "Tổng mét máy QS tháng 4/2026?",
                "Máy hồ Bng tháng này chạy được bao nhiêu mét?",
                "Máy direct tháng này sản xuất bao nhiêu mét beam?",
                "Máy sectional tháng 4 chạy được bao nhiêu mét?",
                "Sản lượng máy winder tuần trước là bao nhiêu?",
                # 🚀 Tốc độ & thời gian
                "So sánh tốc độ thực tế và mục tiêu máy Karlmayer",
                "Thời gian chạy trung bình mỗi lần máy hồ là bao lâu?",
                "Tốc độ máy direct tháng này so tháng trước?",
                "Máy hồ dừng lâu nhất tháng 4 do nguyên nhân gì?",
                # 🏷️ Mã hàng sizing
                "Máy hồ đang chạy mã hàng gì nhiều nhất?",
                "Mã hàng nào được sizing nhiều nhất tháng 4?",
                "Máy sectional tháng 4 chủ yếu chạy mã hàng gì?",
                # 📈 Xu hướng
                "Xu hướng sản lượng máy hồ 10 ngày gần nhất?",
                "Máy sizing nào hiệu suất tăng đều nhất 3 tháng qua?",
                # 🔗 Liên kết sizing → weaving
                "Tháng 4 máy hồ xuất bao nhiêu beam sang weaving?",
                "Beam sizing xuất sang weaving tháng 3 bao nhiêu?",
            ],
            "📊 So sánh / Phân tích": [
                # 🔄 So sánh xưởng — hiệu suất
                "So sánh hiệu suất xưởng 1 và xưởng 2 tháng 3",
                "Xưởng nào sản lượng cao nhất tháng 4/2026?",
                "Hiệu suất YD so với PD xưởng 1 tháng 4?",
                "Xưởng 1 và xưởng 3 hiệu suất tháng nào cao hơn?",
                "So sánh hiệu suất ca A và ca B xưởng 2 tháng 4?",
                # 📅 So sánh kỳ
                "So sánh sản lượng tháng 3, 4, 5/2026 xưởng 2?",
                "So sánh hiệu suất tuần này và tuần trước xưởng 3?",
                "Sản lượng tháng 4 so với tháng 3 tăng hay giảm?",
                "Quý 1/2026 xưởng nào đạt sản lượng cao nhất?",
                "Tháng 5/2026 sản lượng toàn nhà máy bao nhiêu?",
                # 🏷️ Phân tích mã hàng
                "Mã hàng nào chạy nhiều nhất toàn nhà máy tháng này?",
                "Top 5 mã hàng YD sản lượng cao nhất tháng 4?",
                "Mã hàng nào chạy ở cả 3 xưởng tháng 4?",
                "Mã hàng nào bị dừng nhiều nhất tháng 3?",
                "Mã CS 32 tháng 4 chạy xưởng nào nhiều nhất?",
                # 🔗 Liên kết đa cụm
                "Từ kho sợi đến xưởng dệt tháng 4 tổng bao nhiêu kg?",
                "Beam sizing xuất và weaving sử dụng tháng 4 có khớp không?",
                "Máy hồ cung cấp beam cho weaving tháng 3 đủ không?",
                # 📈 Phân tích nâng cao
                "Máy nào xưởng 1 hiệu suất cải thiện nhiều nhất tháng 4?",
                "Máy nào dưới chuẩn (<60%) liên tục 3 tháng?",
                "Xưởng 2 sản lượng trung bình mỗi ngày tháng 4 là bao nhiêu?",
                "Ngày nào toàn nhà máy sản lượng thấp nhất tháng 4?",
            ],
        }
        # ── Question status init ─────────────────────────────────────────
        if "q_status" not in st.session_state:
            st.session_state.q_status = {}

        with st.expander("💡 Gợi ý câu hỏi", expanded=False):
            _n_ok  = sum(1 for v in st.session_state.q_status.values() if v is True)
            _n_bad = sum(1 for v in st.session_state.q_status.values() if v is False)
            _n_total = sum(len(v) for v in _QUESTIONS.values())
            _c1, _c2, _c3 = st.columns(3)
            with _c1: st.caption(f"✅ Đã hoạt động: **{_n_ok}**")
            with _c2: st.caption(f"❌ Cần sửa: **{_n_bad}**")
            with _c3: st.caption(f"⬜ Chưa test: **{_n_total - _n_ok - _n_bad}**")
            st.caption("👆 Click nút trạng thái để đánh dấu: ⬜ → ✅ → ❌ → ⬜")

            _qtabs = st.tabs(list(_QUESTIONS.keys()))
            for _ti, (_cat, _qs) in enumerate(zip(list(_QUESTIONS.keys()), _QUESTIONS.values())):
                with _qtabs[_ti]:
                    _cat_ok  = sum(1 for q in _qs if st.session_state.q_status.get(q) is True)
                    _cat_bad = sum(1 for q in _qs if st.session_state.q_status.get(q) is False)
                    st.caption(f"✅ {_cat_ok} hoạt động | ❌ {_cat_bad} cần sửa | ⬜ {len(_qs)-_cat_ok-_cat_bad} chưa test")
                    _cols_q = st.columns(2)
                    for _qi, _q in enumerate(_qs):
                        with _cols_q[_qi % 2]:
                            _qst = st.session_state.q_status.get(_q)
                            _icon = "✅" if _qst is True else ("❌" if _qst is False else "⬜")
                            _col_btn, _col_cp, _col_pin = st.columns([7, 1, 1])
                            with _col_btn:
                                if st.button(f"{_icon} {_q}", key=f"q_{_ti}_{_qi}",
                                             use_container_width=True, type="secondary"):
                                    st.session_state.suggested_q = _q
                                    st.rerun()
                            with _col_cp:
                                st.markdown(
                                    f'''<button onclick="navigator.clipboard.writeText('{_q.replace("'", "\\'")}')"
                                    style="background:none;border:1px solid #E4E7EE;border-radius:6px;
                                           cursor:pointer;padding:6px 8px;font-size:13px;
                                           height:38px;width:100%;color:#8B92A5"
                                    title="Copy">📋</button>''',
                                    unsafe_allow_html=True
                                )
                            with _col_pin:
                                if st.button(_icon, key=f"pin_{_ti}_{_qi}",
                                             help="Đánh dấu trạng thái",
                                             use_container_width=True):
                                    if _qst is None:
                                        st.session_state.q_status[_q] = True
                                    elif _qst is True:
                                        st.session_state.q_status[_q] = False
                                    else:
                                        st.session_state.q_status[_q] = None
                                    st.rerun()

        # ────────────────────────────────────────────────────────────────

        chat_container = st.container(height=400)
        if "messages" not in st.session_state: st.session_state.messages = []
        
        with chat_container:
            for msg in st.session_state.messages:
                with st.chat_message(msg["role"]): st.markdown(msg["content"])
                    
        # ── Suggested question: keep in state until form submits ──
        _pending_q = st.session_state.get("suggested_q", "")
        _auto_submit = bool(_pending_q)  # auto-submit when suggestion clicked

        with st.form(key="ai_chat_form", clear_on_submit=True):
            col_input, col_btn = st.columns([9, 1])
            with col_input:
                user_prompt = st.text_input(
                    "Hỏi AI", value=_pending_q,
                    label_visibility="collapsed",
                    placeholder="Hỏi về sản lượng, hiệu suất, beam...",
                    key="chat_input_field"
                )
            with col_btn:
                submit_ai = st.form_submit_button("Gửi 🚀", type="primary")

        # Clear pending after form rendered
        if _pending_q:
            st.session_state.suggested_q = ""

        # Use pending_q directly if auto-submit
        _final_prompt = user_prompt or (_pending_q if _auto_submit else "")
        submit_ai = submit_ai or _auto_submit

        if submit_ai and _final_prompt:
            user_prompt = _final_prompt
            st.session_state.messages.append({"role": "user", "content": user_prompt})
            with chat_container:
                st.chat_message("user").markdown(user_prompt)
                with st.chat_message("assistant"):
                    message_placeholder = st.empty() 
                    try:
                        # CHẠY ĐỘNG CƠ LẬT TRANG (YIELD DUMMY CHUNK + TEXT CHUNK)
                        response_stream = process_ai_chat(user_prompt, st.session_state.messages, selected_node, model_name or "qwen2.5:3b", df_cluster, current_view_date)
                        
                        final_text = ""
                        for chunk in response_stream:
                            try:
                                if chunk.text:
                                    final_text += chunk.text
                                    message_placeholder.markdown(final_text + "▌") 
                            except: continue
                        message_placeholder.markdown(final_text)

                    except Exception as e:
                        message_placeholder.markdown(f"Lỗi gọi AI: {str(e)}")

            st.session_state.messages.append({"role": "assistant", "content": final_text})

# ==========================================
# 4. TRANG 2: QUẢN LÝ DỮ LIỆU (CẤU HÌNH TỰ ĐỘNG)
# ==========================================
elif menu_selection == "📥 Quản lý Dữ liệu":
    st.title("QUẢN LÝ DỮ LIỆU KHO & SẢN XUẤT")


    # ══════════════════════════════════════════════════════════════
    # ĐỒNG BỘ THÔNG MINH — chọn nguồn và khoảng thời gian
    # ══════════════════════════════════════════════════════════════
    with st.container(border=True):
        st.markdown("### ⚡ Đồng Bộ Dữ Liệu")

        _ss = st.session_state
        _SOURCES = {
            "weaving":  ("🧵 Xưởng Dệt (TOTAL)",   _ss.get("wv_folder",""),       "thư mục"),
            "sizing":   ("⚙️ Sizing/Sectional",     _ss.get("sizing_folder_path",""), "thư mục"),
            "yarn":     ("📐 Công Thức Sợi",         _ss.get("yarn_folder_saved",""), "thư mục"),
            "beam":     ("🔩 Beam (BEAM_DAT)",       _ss.get("beam_file_path",""),   "file"),
        }

        # Status row
        _cfg_cols = st.columns(4)
        _configured = {}
        for idx, (key, (lbl, path, ptype)) in enumerate(_SOURCES.items()):
            with _cfg_cols[idx]:
                ok = path and (os.path.isdir(path) if ptype=="thư mục" else os.path.isfile(path))
                _configured[key] = ok
                if ok:
                    st.markdown(f'<div style="background:#E6FCF5;border:1px solid #0CA678;border-radius:8px;padding:10px 12px;font-size:13px;font-weight:600;color:#0CA678">✅ {lbl}</div>', unsafe_allow_html=True)
                else:
                    st.markdown(f'<div style="background:#FFF8E1;border:1px solid #F59F00;border-radius:8px;padding:10px 12px;font-size:13px;font-weight:500;color:#B45309">⚠️ {lbl}<br><span style="font-size:11px;font-weight:400">Chưa cấu hình đường dẫn</span></div>', unsafe_allow_html=True)

        st.markdown("<br>", unsafe_allow_html=True)

        # Smart sync controls
        col_sc1, col_sc2, col_sc3 = st.columns([2, 2, 2])
        with col_sc1:
            _sync_mode = st.radio(
                "Chế độ đồng bộ:", 
                ["⚡ Tất cả đã cấu hình", "🎯 Chọn nguồn cụ thể", "📅 Theo thời gian"],
                key="sync_mode_radio", horizontal=False
            )
        with col_sc2:
            if _sync_mode == "🎯 Chọn nguồn cụ thể":
                _selected_sources = st.multiselect(
                    "Chọn nguồn cần đồng bộ:",
                    options=[k for k,v in _configured.items() if v],
                    default=[k for k,v in _configured.items() if v],
                    format_func=lambda k: _SOURCES[k][0],
                    key="sync_sources_select"
                )
            elif _sync_mode == "📅 Theo thời gian":
                import datetime as _dt
                _sync_from = st.date_input("Từ ngày:", value=_dt.date.today()-_dt.timedelta(days=30), key="sync_from_date")
                _sync_to   = st.date_input("Đến ngày:", value=_dt.date.today(), key="sync_to_date")
                _selected_sources = [k for k,v in _configured.items() if v]
                st.caption(f"Chỉ đọc file chứa dữ liệu từ {_sync_from} đến {_sync_to}")
            else:
                _selected_sources = [k for k,v in _configured.items() if v]
        with col_sc3:
            st.markdown("<br>", unsafe_allow_html=True)
            _n_ready = len(_selected_sources) if "_selected_sources" in dir() else len([k for k,v in _configured.items() if v])
            btn_smart_sync = st.button(
                f"⚡ Đồng bộ {_n_ready} nguồn" if _n_ready > 0 else "⚡ Đồng bộ",
                key="btn_master_sync", type="primary",
                use_container_width=True,
                disabled=(_n_ready == 0)
            )

        if btn_smart_sync and _selected_sources:
            _pbar = st.progress(0, text="Đang đồng bộ...")
            _logs = []
            _total = 0
            _n = len(_selected_sources)
            _date_filter = (str(_sync_from), str(_sync_to)) if _sync_mode == "📅 Theo thời gian" and "_sync_from" in dir() else None

            for _si, _src in enumerate(_selected_sources):
                lbl, path, ptype = _SOURCES[_src]
                _pbar.progress(int((_si/_n)*90), text=f"[{_si+1}/{_n}] {lbl}...")
                try:
                    if _src == "weaving":
                        from weaving_pipeline import scan_weaving_folder as _swv2
                        _r = _swv2(path, _ss.get("wv_year","2026"), keyword="TOTAL")
                        _total += _r["total_rows"]
                        _logs.append(f"✅ {lbl}: **{_r['total_rows']}** bản ghi từ **{_r['total_files']}** file")
                    elif _src == "sizing":
                        from sizing_pipeline import scan_sizing_folder_v2 as _ssz2
                        _r = _ssz2(path, _ss.get("sizing_kw",""))
                        _total += _r["total_rows"]
                        _logs.append(f"✅ {lbl}: **{_r['total_rows']}** bản ghi từ **{_r['total_files']}** file")
                    elif _src == "yarn":
                        from yarn_parser import scan_folder as _syn2
                        _r = _syn2(path, _ss.get("yarn_kw_saved",""))
                        _total += _r.get("total",0)
                        _logs.append(f"✅ {lbl}: **{_r.get('total',0)}** bản ghi")
                    elif _src == "beam":
                        from beam_info import import_beam_file as _ibm2
                        _r = _ibm2(path)
                        _total += _r.get("total",0)
                        _logs.append(f"✅ {lbl}: **{_r.get('total',0)}** bản ghi")
                except Exception as _ex:
                    _logs.append(f"❌ {lbl}: {_ex}")

            _pbar.progress(100, text="✅ Hoàn thành!")
            st.success(f"🎉 Đồng bộ xong — tổng **{_total:,} bản ghi** từ **{len(_selected_sources)} nguồn**!")
            for _ll in _logs: st.markdown(f"  {_ll}")
            st.rerun()
    # ══════════════════════════════════════════════════════════════

    with st.container(border=True):
        col_info1, col_info2, col_info3 = st.columns(3)
        with col_info1:
            d_nodes_keys = [
                    "Xưởng Dệt", "Kho Beam Weaving", "Kho Sợi Weaving",
                    "Máy Hồ", "Máy Sectional", "Máy Direct", "Máy Winder", "Máy Suzuki",
                    "Kho Sợi Sizing", "Kho Beam Sizing",
                    "Kho Sợi Tổng", "Xưởng Nhuộm", "Kho Thành Phẩm",
                ]
            cluster_name = st.selectbox("Tên Cụm:", d_nodes_keys, index=9)
        with col_info2:
            st.info("💡 Hệ thống đang chạy cơ chế CỘNG DỒN lịch sử các ngày.")
        with col_info3:
            skip_rows = st.number_input("Bỏ qua dòng đầu (Tính năng rất quan trọng):", value=0, min_value=0)

    # --- SETUP CẤU HÌNH TỰ ĐỘNG VỚI FILE GỐC ---
    st.markdown("### ⚙️ Thiết lập Quét Thư Mục Tự Động (Auto-Sync)")
    # ── ĐỒNG BỘ DỮ LIỆU XƯỞNG DỆT (TOTAL file) ──────────────────
    st.markdown("---")
    with st.expander("🔄 Đồng bộ Dữ liệu Xưởng Dệt (File TOTAL Weaving)", expanded=True):
        # Import weaving_pipeline
        _wv_err = None
        try:
            from weaving_pipeline import import_weaving_total, scan_weaving_folder, init_weaving_table
            init_weaving_table()
        except Exception as _we: _wv_err = str(_we)

        if _wv_err:
            st.error(f"weaving_pipeline.py loi: `{_wv_err}`")
        else:
            # Path persistence
            if "wv_folder" not in st.session_state: st.session_state.wv_folder = ""
            if "wv_year" not in st.session_state: st.session_state.wv_year = "2026"

            col_wv1, col_wv2, col_wv3 = st.columns([3, 1, 1])
            with col_wv1:
                _wv_folder = st.text_input(
                    "📁 Thư mục chứa file TOTAL Weaving:", value=st.session_state.wv_folder,
                    key="wv_folder_input",
                    placeholder="Z:\\2.제직동[WEAVING]\\8. SAN XUAT- 생산\\2026",
                )
                if _wv_folder != st.session_state.wv_folder: st.session_state.wv_folder = _wv_folder
            with col_wv2:
                _wv_year = st.text_input("Năm:", value=st.session_state.wv_year, key="wv_year_input")
                if _wv_year != st.session_state.wv_year: st.session_state.wv_year = _wv_year
            with col_wv3:
                st.markdown("<br>", unsafe_allow_html=True)
                btn_wv_sync = st.button("🔄 Đồng bộ thư mục", key="btn_wv_sync",
                                        type="primary", use_container_width=True,
                                        disabled=(not _wv_folder or not os.path.isdir(_wv_folder)))

            # File uploader
            wv_file = st.file_uploader("📎 Hoặc tải file TOTAL trực tiếp:", type=["xlsx","xls"], key="wv_upload")
            _wv_src = None
            _wv_sheets = ["wea 1","wea 2","wea 3"]
            if wv_file:
                import tempfile as _twv
                _raw_wv = wv_file.read(); wv_file.seek(0)
                with _twv.NamedTemporaryFile(suffix=".xlsx", delete=False) as _tp:
                    _tp.write(_raw_wv); _wv_src = _tp.name
                try:
                    _sh_wv = [s for s in pd.ExcelFile(_wv_src).sheet_names if s in ["wea 1","wea 2","wea 3"]]
                    st.caption(f"Sheets co san: {_sh_wv}")
                    _wv_sheets = st.multiselect("Sheets can nhap:", options=_sh_wv, default=_sh_wv, key="wv_sheets_ms")
                except Exception as _e5: st.warning(str(_e5))

                col_wva, col_wvb = st.columns(2)
                with col_wva:
                    btn_wv_up = st.button("📥 Nhap tu file tai len", key="btn_wv_upload",
                                          type="primary", use_container_width=True,
                                          disabled=(not _wv_src or not _wv_sheets))
                with col_wvb:
                    _wv_year2 = st.text_input("Nam cua file:", value=_wv_year, key="wv_year_upload")
                if btn_wv_up and _wv_src:
                    with st.spinner("Dang nhap..."):
                        _r_wv = import_weaving_total(_wv_src, _wv_sheets, _wv_year2)
                    try: os.unlink(_wv_src)
                    except: pass
                    _detail = " | ".join(f"{k}: {v}" for k,v in _r_wv.items() if k not in ("total","errors") and isinstance(v,int) and v>0)
                    st.success(f"Nhap {_r_wv['total']} ban ghi ({_detail})!")
                    if _r_wv.get("errors"): st.warning("; ".join(_r_wv["errors"][:3]))
                    st.rerun()

            # Folder sync with progress
            if btn_wv_sync and os.path.isdir(_wv_folder):
                _files_wv = [str(p) for p in __import__("pathlib").Path(_wv_folder).rglob("*.xlsx")
                             if "TOTAL" in str(p) and not str(p.name).startswith("~$")]
                if _files_wv:
                    _pb = st.progress(0, text="Dang quet...")
                    _st = st.empty()
                    def _cb_wv(i, total, fname):
                        pct = int((i+1)/total*100) if total else 100
                        _pb.progress(pct, text=f"[{i+1}/{total}] {fname[:50]}")
                        _st.caption(fname)
                    _r5 = scan_weaving_folder(_wv_folder, _wv_year, keyword="TOTAL", progress_cb=_cb_wv)
                    _pb.progress(100, text="Hoan thanh!")
                    st.success(f"Dong bo xong: {_r5['total_rows']} ban ghi tu {_r5['total_files']}/{_r5['total_scanned']} file!")
                    st.rerun()
                else:
                    st.warning("Khong tim thay file TOTAL trong thu muc nay.")

        # Summary table
        try:
            _df_wv = execute_query("SELECT sub_location AS Xuong, date AS Ngay, ten_may AS May, item_id AS Ma_hang, quantity_kg AS KG, hieu_suat_2ca AS HS_2ca, hieu_suat_3ca AS HS_3ca, color AS Color, p_beam_yarn AS Beam_yarn, rpm AS RPM FROM Inventory_Log WHERE cluster_name='Xuong Det' ORDER BY date DESC, sub_location LIMIT 200")
            if not _df_wv.empty:
                st.markdown(f"**{len(_df_wv)} ban ghi** moi nhat trong Inventory_Log (Xuong Det)")
                st.dataframe(_df_wv, use_container_width=True, hide_index=True)
        except: pass
    # ─────────────────────────────────────────────────────────────────

    with st.expander("🧵 Quản lý Công Thức Sợi (Bảng Mẫu)", expanded=False):

        if _yarn_import_error:
            st.error(
                f"⚠️ **yarn_parser.py chưa được tìm thấy trong thư mục chạy app.**\n\n"
                f"Hãy copy file `yarn_parser.py` vào cùng thư mục với `app.py` rồi restart app.\n\n"
                f"Chi tiết lỗi: `{_yarn_import_error}`"
            )
        with st.container(border=True):
            col_y1, col_y2 = st.columns(2)
            # Path persistence
            if "yarn_folder_saved" not in st.session_state:
                st.session_state.yarn_folder_saved = ""
            if "yarn_kw_saved" not in st.session_state:
                st.session_state.yarn_kw_saved = ""

            with col_y1:
                yarn_folder = st.text_input(
                    "📁 Thư mục chứa file Bảng Mẫu:",
                    value=st.session_state.yarn_folder_saved,
                    key="yarn_folder",
                    placeholder="VD: Z:\\VINA_folder\\Bang mau sx moi (chi su dung cai nay)",
                    help="Đường dẫn được lưu lại để đồng bộ nhanh lần sau"
                )
                if yarn_folder != st.session_state.yarn_folder_saved:
                    st.session_state.yarn_folder_saved = yarn_folder
            with col_y2:
                yarn_kw = st.text_input(
                    "🎯 Từ khóa lọc file (để trống = lấy tất cả):",
                    value=st.session_state.yarn_kw_saved,
                    key="yarn_kw",
                    placeholder="VD: NEW, TOWEL..."
                )
                if yarn_kw != st.session_state.yarn_kw_saved:
                    st.session_state.yarn_kw_saved = yarn_kw

            col_ya, col_yb, col_yc = st.columns([2, 2, 3])
            with col_ya:
                yarn_upload = st.file_uploader(
                    "📎 Hoặc tải lên file đơn lẻ:", type=["xlsx","xls"],
                    key="yarn_upload", label_visibility="visible"
                )
            with col_yc:
                # Sync nhanh nếu đã có đường dẫn lưu
                _has_yarn_path = bool(st.session_state.yarn_folder_saved and
                                      os.path.isdir(st.session_state.yarn_folder_saved))
                btn_yarn_scan = st.button(
                    "🔄 Đồng bộ Thư Mục" if _has_yarn_path else "🔄 Quét Thư Mục",
                    key="btn_yarn_scan", type="primary", use_container_width=True,
                    help="Quét và cập nhật toàn bộ file trong thư mục đã lưu"
                )
                btn_yarn_upload = st.button("📤 Lưu File Vừa Tải", key="btn_yarn_upload",
                                             use_container_width=True,
                                             disabled=(yarn_upload is None))

            if btn_yarn_scan and yarn_folder:
                if not os.path.isdir(yarn_folder):
                    st.error(f"❌ Không tìm thấy thư mục: `{yarn_folder}`")
                else:
                    with st.spinner("Đang quét..."):
                        _recs, _files, _errs = scan_folder(yarn_folder, yarn_kw)
                    st.success(f"✅ Đã lưu **{_recs} bản ghi** từ **{_files} file**!")
                    if _errs:
                        with st.expander(f"⚠️ {len(_errs)} file không đọc được"):
                            for e in _errs: st.text(e)
                    st.rerun()

            if btn_yarn_upload and yarn_upload:
                import tempfile
                with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
                    tmp.write(yarn_upload.read())
                    tmp_path = tmp.name
                recs = parse_excel_file(tmp_path)
                # Gán đúng tên file gốc
                for r in recs: r["file_name"] = yarn_upload.name
                n = upsert_records(recs)
                os.unlink(tmp_path)
                st.success(f"✅ Đã lưu **{n} bản ghi** từ `{yarn_upload.name}`!")
                st.rerun()

        # Hiển thị bảng công thức sợi
        try:
            df_yarn = execute_query("""
                SELECT item_name AS 'Tên hàng',
                       ten_may AS 'Loại máy',
                       soi_bong AS 'Sợi bông (loại)', soi_bong_pct AS 'Sợi bông %',
                       soi_ngang AS 'Sợi ngang (loại)', soi_ngang_pct AS 'Sợi ngang %',
                       soi_nen AS 'Sợi nền (loại)', soi_nen_pct AS 'Sợi nền %',
                       soi_border AS 'Sợi border (loại)', soi_border_pct AS 'Sợi border %',
                       file_name AS 'File nguồn'
                FROM Yarn_Formula ORDER BY item_name, ten_may
            """)
            if not df_yarn.empty:
                st.markdown(f"**{len(df_yarn)} bản ghi** trong bảng Công Thức Sợi")
                # Format percentage columns
                for col in ['Sợi bông %', 'Sợi ngang %', 'Sợi nền %', 'Sợi border %']:
                    if col in df_yarn.columns:
                        df_yarn[col] = df_yarn[col].apply(
                            lambda x: f"{x:.2f}%" if pd.notna(x) and x else ""
                        )
                search_yarn = st.text_input("🔍 Tìm mã hàng:", key="search_yarn", placeholder="VD: TOWEL, CM 50S...")
                if search_yarn:
                    mask = df_yarn.apply(lambda r: r.astype(str).str.contains(search_yarn, case=False).any(), axis=1)
                    df_yarn = df_yarn[mask]
                st.dataframe(df_yarn, use_container_width=True, hide_index=True)

                # Delete button
                if st.button("🗑️ Xóa toàn bộ bảng Công Thức Sợi", key="del_yarn"):
                    try:
                        import sqlite3 as _sl3
                        _db_files = glob.glob("*.db") + glob.glob("**/*.db", recursive=True)
                        for _dbf in _db_files:
                            with _sl3.connect(_dbf) as _c:
                                _c.execute("DELETE FROM Yarn_Formula")
                                _c.commit()
                        st.success("✅ Đã xóa toàn bộ bảng Công Thức Sợi!")
                    except Exception as _de:
                        st.error(f"Lỗi xóa: {_de}")
                    st.rerun()
            else:
                st.info("Chưa có dữ liệu. Quét thư mục hoặc tải file để bắt đầu.")
        except Exception as _ye:
            st.warning(f"Bảng Yarn_Formula chưa được khởi tạo: {_ye}")
        # ─────────────────────────────────────────────────────────────────

        # ── THÔNG TIN BEAM (BEAM_DAT) ────────────────────────────────────
        with st.container(border=True):
            if _beam_import_error:
                st.error(f"⚠️ beam_info.py chưa tìm thấy: `{_beam_import_error}`")
            else:
                # ── Lưu đường dẫn file vào session_state để sync lại ──
                if "beam_file_path" not in st.session_state:
                    st.session_state.beam_file_path = ""

                col_b0, col_b1 = st.columns([3, 2])
                with col_b0:
                    _beam_folder = st.text_input(
                        "📁 Đường dẫn thư mục hoặc file BEAM_DAT:",
                        value=st.session_state.beam_file_path,
                        key="beam_folder_input",
                        placeholder="VD: Z:\VINA_folder\BEAM_DAT-THONG_TIN_BEAM_2025.xlsx",
                    )
                    if _beam_folder != st.session_state.beam_file_path:
                        st.session_state.beam_file_path = _beam_folder

                with col_b1:
                    st.markdown("<br>", unsafe_allow_html=True)
                    btn_beam_sync = st.button(
                        "🔄 Đồng bộ từ đường dẫn", key="btn_beam_sync",
                        use_container_width=True,
                        disabled=(not _beam_folder or not os.path.isfile(_beam_folder))
                    )

                # Upload file thủ công
                beam_file = st.file_uploader(
                    "📎 Hoặc tải file trực tiếp:",
                    type=["xlsx","xls"], key="beam_file_upload",
                )

                # Sheet selector — hiển thị khi có file upload
                _beam_sheets = []
                _beam_src_path = None
                if beam_file is not None:
                    import tempfile as _tb
                    _raw = beam_file.read()
                    beam_file.seek(0)
                    with _tb.NamedTemporaryFile(suffix=".xlsx", delete=False) as _tp:
                        _tp.write(_raw); _tp.flush()
                        _tp_path = _tp.name
                    try:
                        _all_sh = pd.ExcelFile(_tp_path).sheet_names
                        _supported = [s for s in ["TOTAL","XUATKHO","YCCB"] if s in _all_sh]
                        _other_sh = len(_all_sh) - len(_supported)
                        st.caption(f"📑 File có **{len(_all_sh)} sheets** ({len(_supported)} hỗ trợ)")
                        if not _supported:
                            st.warning("⚠️ Không tìm thấy sheet TOTAL/XUATKHO/YCCB trong file này!")
                        _beam_sheets = st.multiselect(
                            "Sheets cần nhập:",
                            options=_supported,
                            default=["TOTAL"] if "TOTAL" in _supported else _supported,
                            key="beam_sheets_ms",
                            help="TOTAL = toàn bộ lịch sử beam | XUATKHO = beam xuất kho | YCCB = yêu cầu"
                        )
                        _beam_src_path = _tp_path
                    except Exception as _she:
                        st.warning(f"Không đọc được sheets: {_she}")

                col_ba, col_bb = st.columns(2)
                with col_ba:
                    btn_beam_upload = st.button(
                        "📥 Nhập Beam từ file tải lên", key="btn_beam_import",
                        type="primary", use_container_width=True,
                        disabled=(beam_file is None or not _beam_sheets)
                    )
                with col_bb:
                    if _beam_folder and os.path.isfile(_beam_folder):
                        st.caption(f"✅ Đường dẫn hợp lệ")
                    elif _beam_folder:
                        st.caption("⚠️ Đường dẫn chưa tìm thấy file")

                # ── Xử lý UPLOAD ──
                if btn_beam_upload and beam_file and _beam_src_path and _beam_sheets:
                    with st.spinner("Đang nhập..."):
                        from beam_info import read_xuatkho, read_yccb, upsert_beam_info, upsert_beam_request, init_beam_tables
                        init_beam_tables()
                        _fname = beam_file.name
                        _nx = _ny = 0
                        from beam_info import read_total as _read_total
                        if "TOTAL" in _beam_sheets:
                            _nx += upsert_beam_info(_read_total(_beam_src_path, _fname))
                        if "XUATKHO" in _beam_sheets:
                            _nx += upsert_beam_info(read_xuatkho(_beam_src_path, _fname))
                        if "YCCB" in _beam_sheets:
                            _ny = upsert_beam_request(read_yccb(_beam_src_path, _fname))
                    try: os.unlink(_beam_src_path)
                    except: pass
                    st.success(f"✅ Nhập **{_nx} beam** vào Beam_Info + **{_ny} yêu cầu** vào Beam_Request!")
                    if _beam_folder == "":
                        st.session_state.beam_file_path = ""  # suggest saving path
                    st.rerun()

                # ── Xử lý SYNC từ đường dẫn ──
                if btn_beam_sync and os.path.isfile(_beam_folder):
                    with st.spinner(f"Đang đọc {os.path.basename(_beam_folder)}..."):
                        from beam_info import import_beam_file, init_beam_tables
                        init_beam_tables()
                        _res2 = import_beam_file(_beam_folder)
                    _total_imported = _res2.get('total_sheet', 0) + _res2.get('xuatkho', 0)
                    st.success(f"✅ Đồng bộ xong: **{_total_imported} beam** (TOTAL+XUATKHO) + **{_res2['yccb']} yêu cầu**!")
                    st.rerun()

        # Hiển thị bảng Beam_Info
        try:
            _df_beam = execute_query("""
                SELECT ma_beam AS 'Mã Beam',
                       weaving AS 'Xưởng', so_may AS 'Máy', loai_may AS 'Loại máy',
                       ten_hang AS 'Tên hàng', loai_soi AS 'Loại sợi',
                       so_met AS 'Số mét', so_kg_thuc_te AS 'Kg thực tế',
                       phan_loai AS 'PD/YD',
                       ngay_len_may AS 'Ngày lên máy', ghi_chu AS 'Ghi chú'
                FROM Beam_Info ORDER BY ngay_len_may DESC, weaving, CAST(so_may AS REAL)
            """)
            if not _df_beam.empty:
                st.markdown(f"**{len(_df_beam)} beam** trong bảng Beam_Info (đã lên máy)")
                _col_f1, _col_f2 = st.columns(2)
                with _col_f1:
                    _beam_search = st.text_input("🔍 Tìm beam/tên hàng:", key="beam_search",
                                                  placeholder="VD: 26-3167, SW NEW COLOR...")
                with _col_f2:
                    _beam_wv = st.selectbox("Lọc Xưởng:", ["Tất cả","Weaving 1","Weaving 2","Weaving 3"],
                                             key="beam_wv_filter")
                _df_show = _df_beam.copy()
                if _beam_search:
                    _mask = _df_show.apply(lambda r: r.astype(str).str.contains(_beam_search, case=False).any(), axis=1)
                    _df_show = _df_show[_mask]
                if _beam_wv != "Tất cả":
                    _df_show = _df_show[_df_show["Xưởng"].astype(str).str.contains(_beam_wv.replace("Weaving ",""), na=False)]
                st.dataframe(_df_show, use_container_width=True, hide_index=True)
                if st.button("🗑️ Xóa bảng Beam_Info", key="del_beam"):
                    import sqlite3 as _sl3b
                    for _dbf in glob.glob("*.db"):
                        with _sl3b.connect(_dbf) as _c:
                            _c.execute("DELETE FROM Beam_Info")
                            _c.execute("DELETE FROM Beam_Request")
                            _c.commit()
                    st.success("Đã xóa!"); st.rerun()
            else:
                st.info("Chưa có dữ liệu beam. Tải file BEAM_DAT để bắt đầu.")
        except Exception as _bex:
            st.warning(f"Bảng Beam_Info chưa tồn tại: {_bex}")
        # ─────────────────────────────────────────────────────────────────

    # ── SIZING / SECTIONAL / DIRECT / WINDER ────────────────────────
    st.markdown("---")
    # ── SIZING UI ──────────────────────────────────────────────────────
    st.markdown("---")
    with st.expander("⚙️ Quản lý Dữ liệu Máy Sizing / Sectional / Direct / Winder", expanded=False):
        if _sizing_import_error:
            st.error(f"sizing_pipeline.py lỗi: `{_sizing_import_error}`")
        else:
            # Path persistence
            if "sizing_folder_path" not in st.session_state:
                st.session_state.sizing_folder_path = ""
            if "sizing_kw" not in st.session_state:
                st.session_state.sizing_kw = ""

            col_szA, col_szB = st.columns([3,2])
            with col_szA:
                _sz_folder = st.text_input(
                    "📁 Thư mục gốc (chứa các năm/tháng/file):",
                    value=st.session_state.sizing_folder_path,
                    key="sz_folder_input",
                    placeholder="VD: Z:\\1.준비동[SIZING]\\3. 데이터(DATA)\\1. SẢN LƯỢNG",
                    help="Hệ thống sẽ tự quét đệ quy: Năm → Tháng → File"
                )
                if _sz_folder != st.session_state.sizing_folder_path:
                    st.session_state.sizing_folder_path = _sz_folder
            with col_szB:
                _sz_kw = st.text_input("🎯 Từ khóa lọc file:", value=st.session_state.sizing_kw,
                                        key="sz_kw_input", placeholder="VD: HIỆU SUẤT")
                if _sz_kw != st.session_state.sizing_kw:
                    st.session_state.sizing_kw = _sz_kw

            # Preview file list
            if _sz_folder and os.path.isdir(_sz_folder):
                from sizing_pipeline import list_files_in_folder
                _preview_files = list_files_in_folder(_sz_folder, _sz_kw)
                _n_files = len(_preview_files)
                if _n_files > 0:
                    st.success(f"✅ Tìm thấy **{_n_files} file** trong thư mục (đệ quy qua các năm/tháng)")
                    with st.expander(f"📂 Xem danh sách {min(_n_files,20)} file đầu tiên"):
                        for fp in _preview_files[:20]:
                            # Show relative path from root
                            try:
                                rel = os.path.relpath(fp, _sz_folder)
                            except:
                                rel = os.path.basename(fp)
                            st.text(f"  📄 {rel}")
                        if _n_files > 20:
                            st.caption(f"... và {_n_files-20} file khác")
                else:
                    st.warning("Không tìm thấy file xlsx/xls nào trong thư mục này")
            elif _sz_folder:
                st.warning("⚠️ Thư mục chưa tìm thấy")

            col_szC, col_szD = st.columns(2)
            with col_szC:
                btn_sz_scan = st.button("🔄 Đồng bộ toàn bộ thư mục", key="btn_sz_scan",
                                         type="primary", use_container_width=True,
                                         disabled=(not _sz_folder or not os.path.isdir(_sz_folder)))
            with col_szD:
                sz_file = st.file_uploader("📎 Hoặc tải 1 file:", type=["xlsx","xls"], key="sz_upload_single")

            # Upload 1 file
            _sz_src = None
            if sz_file:
                import tempfile as _tsz2
                _raw2 = sz_file.read(); sz_file.seek(0)
                with _tsz2.NamedTemporaryFile(suffix=".xlsx", delete=False) as _tp2:
                    _tp2.write(_raw2); _tp2_path = _tp2.name
                try:
                    _all_sh2 = pd.ExcelFile(_tp2_path).sheet_names
                    _avail2  = [s for s in ["MÁY HỒ","MÁY SEC","MÁY QS","SUZUKI","WINDER"] if s in _all_sh2]
                    st.caption(f"📑 {len(_all_sh2)} sheets | {len(_avail2)} hỗ trợ")
                    _sz_sheets2 = st.multiselect("Sheets:", options=_avail2, default=_avail2, key="sz_sheets_single")
                    _sz_src = _tp2_path
                except Exception as _e3: st.warning(str(_e3))
                btn_sz_upload = st.button("📥 Nhập file này", key="btn_sz_upload_single",
                                           type="primary", use_container_width=True,
                                           disabled=(not _sz_src))
                if btn_sz_upload and _sz_src:
                    with st.spinner("Đang nhập..."):
                        _r3 = detect_and_import(_sz_src)
                    try: os.unlink(_sz_src)
                    except: pass
                    _ft = _r3.get("file_type","")
                    _ft_lbl = "Preparation_Product_Report" if _ft=="prep_report" else "FILE_HIỆU_SUẤT_TỔNG"
                    st.success(f"✅ Nhập **{_r3['total']} bản ghi** từ `{_ft_lbl}`!")
                    st.rerun()

            # Folder scan with progress
            if btn_sz_scan:
                from sizing_pipeline import scan_sizing_folder as _scan_sz
                _pbar = st.progress(0, text="Đang quét...")
                _status = st.empty()
                def _cb(i, total, fname):
                    pct = int((i+1)/total*100) if total else 100
                    _pbar.progress(pct, text=f"[{i+1}/{total}] {fname[:50]}")
                    _status.caption(f"Đang xử lý: {fname}")
                with st.spinner(""):
                    _r4 = scan_sizing_folder_v2(_sz_folder, _sz_kw, progress_cb=_cb)
                _pbar.progress(100, text="✅ Hoàn thành!")
                _det = ""
                if _r4.get("prep_report"): _det += f" | Prep Report: {_r4['prep_report']}"
                if _r4.get("hieu_suat"):   _det += f" | Hiệu Suất: {_r4['hieu_suat']}"
                st.success(
                    f"🎉 Đã đồng bộ **{_r4['total_rows']} bản ghi** "
                    f"từ **{_r4['total_files']}/{_r4['total_scanned']} file**{_det}!"
                )
                if _r4.get("errors"):
                    with st.expander(f"⚠️ {len(_r4['errors'])} lỗi"):
                        for _e4 in _r4["errors"][:10]: st.text(_e4)
                st.rerun()

        # Preview table
        try:
            _df_sz = execute_query("""SELECT machine_type AS May, date AS Ngay,
                ten_may AS Ten_may, ca AS Ca, loai_soi AS Loai_soi,
                toc_do_thuc_te AS Toc_do, sl_thuc_te_mtr AS SL_MTR, sl_kg AS SL_KG,
                hieu_suat_pct AS HS_pct, gio_bat_dau AS Bat_dau,
                gio_ket_thuc AS Ket_thuc, thoi_gian_phut AS TG_phut
                FROM Sizing_Log ORDER BY date DESC, machine_type LIMIT 500""")
            if not _df_sz.empty:
                st.markdown(f"**{len(_df_sz)} bản ghi** trong Sizing_Log")
                col_sf1, col_sf2 = st.columns([5,1])
                with col_sf1:
                    _sz_q = st.text_input("🔍 Lọc:", key="sz_filter", placeholder="Tên máy, loại sợi...")
                with col_sf2:
                    st.markdown("<br>", unsafe_allow_html=True)
                    if st.button("🗑️ Xóa tất cả", key="sz_del"):
                        import sqlite3 as _sl4
                        for _db4 in glob.glob("*.db"):
                            with _sl4.connect(_db4) as _c4: _c4.execute("DELETE FROM Sizing_Log"); _c4.commit()
                        st.success("Đã xóa!"); st.rerun()
                if _sz_q:
                    _msk = _df_sz.apply(lambda r: r.astype(str).str.contains(_sz_q, case=False).any(), axis=1)
                    _df_sz = _df_sz[_msk]
                st.dataframe(_df_sz, use_container_width=True, hide_index=True)
            else:
                st.info("Chưa có dữ liệu. Quét thư mục hoặc tải file để bắt đầu.")
        except Exception as _sxe:
            st.warning(f"Sizing_Log chưa tồn tại: {_sxe}")
    # ─────────────────────────────────────────────────────────────────

    # ═══════════════════════════════════════════════════════════════════
    # NHẬT KÝ NỐI SỢI — Parse & 5-file Batch Merge
    # ═══════════════════════════════════════════════════════════════════
    with st.expander("📖 Nhật Ký Nối Sợi & Merge 5 File Báo Cáo", expanded=False):
        st.caption("Upload file nhật ký nối sợi → parse timing từng máy. Hoặc upload 5 file (3 sản lượng + 2 nhật ký) để merge và tải về.")

        try:
            from noi_soi_parser import parse_file as _nsparse, get_available_sheets as _nssheets, merge_into_analysis as _nsmerge
            _ns_ok = True
        except ImportError as _nse:
            st.error(f"noi_soi_parser.py chưa tìm thấy: {_nse}")
            _ns_ok = False

        st.caption(
            "Upload **5 file**: 3 file sản lượng (có sheet *analysis*) + 2 nhật ký nối sợi "
            "→ hệ thống tự nhận diện → merge timing → **tải về 1 file Excel với 3 sheet** (mỗi sheet = 1 xưởng)"
        )

        _b_files = st.file_uploader(
            "Upload 5 file (kéo thả hoặc Browse):",
            type=["xlsx","xls"], accept_multiple_files=True, key="batch_5f",
            help="3 file sản lượng (tên chứa 1동/2동/3동 hoặc wea1/wea2/wea3) + 2 nhật ký (tên chứa nhật ký/타이밍)"
        )
        _b_date = st.text_input(
            "Ngày cần đọc từ nhật ký (VD: 28-05 hoặc 28-5):",
            key="batch_d", placeholder="28-05"
        )

        if _b_files:
            # Auto-classify
            _prod_map = {}   # {weaving_name: file}
            _timing_f = []

            for _bf in _b_files:
                fn = _bf.name
                if any(k in fn.lower() for k in ['nhật ký','nhat ky','타이밍','noi soi','timing']):
                    _timing_f.append(_bf)
                else:
                    fn_l = fn.lower()
                    if '1동' in fn_l or 'wea 1' in fn_l or 'wea1' in fn_l or '제직_1' in fn_l:
                        _prod_map['Weaving 1'] = _bf
                    elif '2동' in fn_l or 'wea 2' in fn_l or 'wea2' in fn_l or '제직_2' in fn_l:
                        _prod_map['Weaving 2'] = _bf
                    elif '3동' in fn_l or 'wea 3' in fn_l or 'wea3' in fn_l or '제직_3' in fn_l:
                        _prod_map['Weaving 3'] = _bf
                    else:
                        # Unknown — add as extra prod file
                        _prod_map[f"Xưởng ({fn[:15]})"] = _bf

            # Show classification
            _cl1, _cl2 = st.columns(2)
            with _cl1:
                st.markdown("**📊 File sản lượng phát hiện:**")
                if _prod_map:
                    for wn, wf in _prod_map.items():
                        st.caption(f"  ✅ **{wn}**: `{wf.name[:40]}`")
                else:
                    st.caption("  ⚠️ Chưa nhận diện được file sản lượng")
            with _cl2:
                st.markdown("**📖 File nhật ký nối sợi:**")
                if _timing_f:
                    for tf in _timing_f:
                        st.caption(f"  ✅ `{tf.name[:40]}`")
                else:
                    st.caption("  ⚠️ Chưa tìm thấy file nhật ký")

            if len(_prod_map) == 0:
                st.warning("Không nhận diện được file sản lượng. Kiểm tra tên file có chứa: 1동/2동/3동 hoặc wea1/wea2/wea3.")
            elif not _timing_f:
                st.warning("Không tìm thấy file nhật ký. Tên file phải chứa: nhật ký, 타이밍, noi soi.")
            elif not _b_date:
                st.info("Nhập ngày cần merge để bắt đầu.")
            else:
                _btn_merge = st.button(
                    f"🔀 Merge & Tạo file {len(_prod_map)} sheet → Tải về",
                    key="btn_bmerge", type="primary", use_container_width=True
                )
                if _btn_merge:
                    with st.spinner("Đang parse nhật ký và merge timing..."):
                        # 1. Parse all timing files
                        _trecs = []
                        for _tf in _timing_f:
                            _tf.seek(0)
                            _tr = _nsparse(_tf, target_date=_b_date, filename=_tf.name)
                            _tf.seek(0)
                            _trecs.extend(_tr.get("records", []))
                            if _tr.get("sheet_used"):
                                st.caption(f"📖 {_tf.name[:35]} → sheet `{_tr['sheet_used']}` → {len(_tr.get('records',[]))} máy")

                        # 2. Build files_info list
                        _files_info = []
                        for _wname, _wf in _prod_map.items():
                            _wf.seek(0)
                            _files_info.append({
                                'bytes': _wf.read(),
                                'filename': _wf.name,
                                'weaving': _wname,
                            })
                            _wf.seek(0)

                        # 3. Create merged 3-sheet output
                        try:
                            from noi_soi_parser import create_merged_output as _cmo
                            _result_bytes, _n_filled = _cmo(_files_info, _trecs, _b_date)
                            _out_fname = f"Merged_Analysis_{_b_date.replace('/','_')}.xlsx"
                            st.success(
                                f"✅ Tạo xong! **{len(_files_info)} sheet** | "
                                f"**{_n_filled} ô timing** được điền từ nhật ký"
                            )
                            st.download_button(
                                label=f"⬇️ Tải về: {_out_fname}",
                                data=_result_bytes,
                                file_name=_out_fname,
                                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                key="dl_merged_final",
                                use_container_width=True
                            )
                            st.caption(
                                f"File gồm {len(_files_info)} sheet: "
                                + " | ".join(f"'{i['weaving']} ({_b_date})'" for i in _files_info)
                            )
                        except Exception as _em:
                            st.error(f"Lỗi tạo file merge: {_em}")

    st.divider()
    st.divider()
    st.subheader("🗄️ Bảng Dữ Liệu SQL Gốc (Xem/Sửa/Xóa)")
    try:
        df_db_raw = execute_query("SELECT * FROM Inventory_Log")
        if not df_db_raw.empty:
            # ── Tab between Inventory_Log and Sizing_Log ──
            _sql_tab1, _sql_tab2 = st.tabs(["🧵 Xưởng Dệt / Kho (Inventory_Log)", "⚙️ Sizing / Sectional / Winder (Sizing_Log)"])

            with _sql_tab1:
                filter_cluster = st.selectbox("👉 Lọc theo Cụm/Kho:", ["-- Tất cả --"] + [c for c in df_db_raw["cluster_name"].dropna().unique()], key="sql_cluster_filter")
            if filter_cluster != "-- Tất cả --":
                df_filtered = get_clean_cluster_data(df_db_raw[df_db_raw['cluster_name'] == filter_cluster], filter_cluster)
                
                # KHỬ TRÙNG LẶP TRƯỚC KHI HIỂN THỊ
                df_filtered = df_filtered.loc[:, ~df_filtered.columns.duplicated()]
                
                loc_filter = "Tất cả"
                if filter_cluster == "Xưởng Dệt":
                    st.markdown("### 🎛️ Bộ lọc Khu vực Xưởng Dệt")
                    loc_filter = st.radio("📍 Chọn khu vực hiển thị dữ liệu:", ["Tất cả", "Weaving 1", "Weaving 2", "Weaving 3"], horizontal=True)
                    
                    if loc_filter != "Tất cả":
                        def match_zone_val(loc_str, target):
                            ls = str(loc_str).lower().strip()
                            if target == "Weaving 1" and ls in ['1', '1.0', 'wea 1', 'wea1', 'weaving 1', 'w1']: return True
                            if target == "Weaving 2" and ls in ['2', '2.0', 'wea 2', 'wea2', 'weaving 2', 'w2']: return True
                            if target == "Weaving 3" and ls in ['3', '3.0', 'wea 3', 'wea3', 'weaving 3', 'w3']: return True
                            if target.lower() in ls: return True
                            return False
                        
                        loc_series = df_filtered['sub_location'] if 'sub_location' in df_filtered.columns else pd.Series(index=df_filtered.index, data="")
                        mask = loc_series.astype(str).apply(lambda x: match_zone_val(x, loc_filter))
                        df_to_edit = df_filtered[mask].copy()
                    else:
                        df_to_edit = df_filtered.copy()
                else:
                    df_to_edit = df_filtered.copy()
                
                df_to_edit = df_to_edit.astype(str).replace(['nan', 'None', '<NA>', 'NaN'], '')
                
                col_config = {}
                if "id" in df_to_edit.columns: col_config["id"] = None
                if "cluster_name" in df_to_edit.columns: col_config["cluster_name"] = None
                
                dynamic_editor_key = f"db_editor_{filter_cluster}_{loc_filter}"
                
                # CHỐNG LỖI KHI BẢNG ĐỂ EDIT RỖNG
                if not df_to_edit.empty:
                    edited_df = st.data_editor(df_to_edit, num_rows="dynamic", use_container_width=True, key=dynamic_editor_key, column_config=col_config)
                else:
                    st.info("Không có dữ liệu khớp với bộ lọc để chỉnh sửa.")
                    edited_df = pd.DataFrame()
                
                c1, c2 = st.columns(2)
                with c1:
                    if st.button("💾 Lưu các chỉnh sửa trên bảng", type="primary", key=f"btn_save_{filter_cluster}_{loc_filter}"):
                        
                        original_ids = [int(i) for i in df_to_edit.get('id', pd.Series(dtype=str)).replace('', pd.NA).dropna().tolist() if str(i).isdigit()]
                        
                        if original_ids:
                            id_placeholders = ",".join("?" * len(original_ids))
                            run_db_command(f"DELETE FROM Inventory_Log WHERE id IN ({id_placeholders})", tuple(original_ids))
                        elif filter_cluster != "Xưởng Dệt" or loc_filter == "Tất cả":
                            run_db_command("DELETE FROM Inventory_Log WHERE cluster_name = ?", (filter_cluster,))
                        
                        if not edited_df.empty:
                            edited_df = edited_df.replace('', pd.NA).dropna(how='all')
                            if not edited_df.empty:
                                if 'id' in edited_df.columns: edited_df = edited_df.drop(columns=['id'])
                                edited_df['cluster_name'] = filter_cluster
                                insert_data(edited_df)
                        st.rerun() 
                with c2:
                    if st.button("🗑️ Xóa SẠCH dữ liệu đang hiển thị", key=f"btn_del_{filter_cluster}_{loc_filter}"):
                        original_ids = [int(i) for i in df_to_edit.get('id', pd.Series(dtype=str)).replace('', pd.NA).dropna().tolist() if str(i).isdigit()]
                        if original_ids:
                            id_placeholders = ",".join("?" * len(original_ids))
                            run_db_command(f"DELETE FROM Inventory_Log WHERE id IN ({id_placeholders})", tuple(original_ids))
                        elif filter_cluster != "Xưởng Dệt" or loc_filter == "Tất cả":
                            run_db_command("DELETE FROM Inventory_Log WHERE cluster_name = ?", (filter_cluster,))
                        st.rerun() 
            else:
                df_display = df_db_raw.drop(columns=['id', 'cluster_name'], errors='ignore')
                df_display = df_display.loc[:, ~df_display.columns.duplicated()]
                df_display = df_display.astype(str).replace(r'^\s*$', pd.NA, regex=True).dropna(axis=1, how='all').fillna("")
                
                # CHỐNG LỖI KHI BẢNG GỐC RỖNG
                if not df_display.empty:
                    st.dataframe(df_display.sort_index(ascending=False), use_container_width=True)
                else:
                    st.info("Bảng trống.")
                st.info(f"Tổng cộng: **{len(df_db_raw)}** dòng.")
        else:
            st.info("📭 Database trống.")

            with _sql_tab2:
                try:
                    _df_sz2 = execute_query("""
                        SELECT machine_type AS May, date AS Ngay, ten_may AS Ten_may,
                               ca AS Ca, loai_soi AS Loai_soi, pd_yd AS PD_YD,
                               toc_do_muc_tieu AS Toc_do_MT, toc_do_thuc_te AS Toc_do_TT,
                               sl_thuc_te_mtr AS SL_MTR, sl_kg AS SL_KG,
                               hieu_suat_pct AS HS_pct,
                               gio_bat_dau AS Bat_dau, gio_ket_thuc AS Ket_thuc,
                               thoi_gian_phut AS TG_phut, ghi_chu AS Ghi_chu
                        FROM Sizing_Log ORDER BY date DESC, machine_type, ten_may
                    """)
                    if not _df_sz2.empty:
                        _sz_type_filter = st.selectbox("Lọc loại máy:",
                            ["-- Tất cả --"] + list(_df_sz2["May"].dropna().unique()),
                            key="sz_type_filter")
                        _sz_search2 = st.text_input("🔍 Tìm:", key="sz_search2",
                            placeholder="Tên máy, loại sợi...")
                        _df_sz_show = _df_sz2.copy()
                        if _sz_type_filter != "-- Tất cả --":
                            _df_sz_show = _df_sz_show[_df_sz_show["May"] == _sz_type_filter]
                        if _sz_search2:
                            _msk = _df_sz_show.apply(lambda r: r.astype(str).str.contains(_sz_search2, case=False).any(), axis=1)
                            _df_sz_show = _df_sz_show[_msk]
                        st.markdown(f"**{len(_df_sz_show)} bản ghi** | {_sz_type_filter}")
                        st.dataframe(_df_sz_show, use_container_width=True, hide_index=True)
                    else:
                        st.info("Chưa có dữ liệu trong Sizing_Log. Vào Quản lý Dữ liệu → đồng bộ file hiệu suất.")
                except Exception as _sze2:
                    st.warning(f"Sizing_Log chưa tồn tại: {_sze2}")
    except Exception as e:
        st.warning(f"Lỗi Database: {str(e)}. Hãy thử F5 lại trang.")